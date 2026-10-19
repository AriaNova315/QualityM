from __future__ import annotations

import asyncio
import time
from contextlib import asynccontextmanager
from pathlib import Path

import httpx
import pytest
from openpyxl import load_workbook

import app.application.template_service as template_service
from app.application.template_service import (
    TemplateServiceError,
    checklist_library,
    commit_import_preview,
    copy_issue,
    create_import_preview,
    create_issue,
    create_template_entity,
    flat_template_rows,
    restore_builtin_entity,
    update_issue,
    update_template_entity,
)
from app.database import initialize_database
from app.excel.template_library import HEADERS, SHEET_NAME, build_template_library, read_template_library
from app.main import create_app


def run(coro): return asyncio.run(coro)


def test_template_library_excel_round_trip_and_hierarchy(test_settings, tmp_path: Path) -> None:
    run(initialize_database(test_settings.database_path,test_settings.migrations_dir,test_settings.bootstrap_seed_path,test_settings.seed_schema_path))
    rows=run(flat_template_rows(test_settings.database_path)); path=build_template_library(rows,tmp_path/"library.xlsx")
    book=load_workbook(path,read_only=True)
    try: assert book.sheetnames == [SHEET_NAME] and tuple(cell.value for cell in book[SHEET_NAME][1]) == HEADERS
    finally: book.close()
    assert read_template_library(path)[0]["issue_key"] == rows[0]["issue_key"]
    library=run(checklist_library(test_settings.database_path)); assert len(library) == 25 and library[0]["groups"]


def test_issue_update_ignores_removed_default_concept(test_settings) -> None:
    run(initialize_database(test_settings.database_path,test_settings.migrations_dir,test_settings.bootstrap_seed_path,test_settings.seed_schema_path))
    rows=run(flat_template_rows(test_settings.database_path)); target=rows[0]
    import sqlite3
    with sqlite3.connect(test_settings.database_path) as c:
        issue_id=c.execute("SELECT id FROM issue_templates WHERE stable_key=?",(target["issue_key"],)).fetchone()[0]
    updated=run(update_issue(test_settings.database_path,issue_id,{"title":"已编辑模板","is_default":True}))
    assert updated["title"] == "已编辑模板" and updated["user_modified"] == 1
    assert updated["is_default"] == 0 and updated["revision"] == 2


def test_template_import_preview_detects_changes_without_writing(test_settings, tmp_path: Path) -> None:
    run(initialize_database(test_settings.database_path,test_settings.migrations_dir,test_settings.bootstrap_seed_path,test_settings.seed_schema_path))
    rows=run(flat_template_rows(test_settings.database_path)); changed=dict(rows[0]); changed["title"]="Excel 修改标题"
    added=dict(changed); added["issue_key"]=f"{changed['issue_key']}.additional"; added["title"]="Excel 新增问题"; added["is_default"]=False; added["issue_order"]=99
    path=build_template_library((changed,added,*rows[1:]),tmp_path/"changed.xlsx")
    preview=run(create_import_preview(test_settings.database_path,path.read_bytes()))
    assert preview["preview_token"] and {change["kind"] for change in preview["changes"]} == {"modified","added"}
    assert run(flat_template_rows(test_settings.database_path))[0]["title"] != "Excel 修改标题"
    committed=run(commit_import_preview(test_settings.database_path,preview["preview_token"]))
    assert committed["applied_changes"] == 2
    assert run(flat_template_rows(test_settings.database_path))[0]["title"] == "Excel 修改标题"
    assert any(row["issue_key"] == added["issue_key"] for row in run(flat_template_rows(test_settings.database_path)))


def test_crud_versions_defaults_disable_limits_and_builtin_restore(test_settings) -> None:
    run(initialize_database(test_settings.database_path,test_settings.migrations_dir,test_settings.bootstrap_seed_path,test_settings.seed_schema_path))
    library = run(checklist_library(test_settings.database_path))
    builtin_item = library[0]["groups"][0]["items"][0]
    original_content = builtin_item["content"]
    edited = run(update_template_entity(test_settings.database_path,"items",builtin_item["id"],{"content":"用户修改内容","display_order":99}))
    assert edited["id"] == builtin_item["id"] and edited["stable_key"] == builtin_item["stable_key"]
    assert edited["version"].endswith("+user.1") and edited["user_modified"] == 1
    restored = run(restore_builtin_entity(test_settings.database_path,test_settings.bootstrap_seed_path,"items",builtin_item["id"]))
    assert restored["content"] == original_content and restored["user_modified"] == 0

    checklist = run(create_template_entity(test_settings.database_path,"checklists",None,{"stable_key":"user.checklist","name":"用户检查单","audit_kind":"product","display_order":99}))
    group = run(create_template_entity(test_settings.database_path,"groups",checklist["id"],{"stable_key":"user.group","name":"用户分组","display_order":1}))
    item = run(create_template_entity(test_settings.database_path,"items",group["id"],{"stable_key":"user.item","display_no":"1","content":"用户检查项","starred":False,"display_order":1}))
    assert checklist["enabled"] == group["enabled"] == item["enabled"] == 0
    assert next(value for value in run(checklist_library(test_settings.database_path)) if value["id"] == checklist["id"])["groups"][0]["items"][0]["id"] == item["id"]
    first = run(create_issue(test_settings.database_path,item["id"],{"stable_key":"user.issue.first","title":"问题一","description":"描述一","corrective_suggestion":"建议一","default_severity":"一般","is_default":True}))
    second = run(create_issue(test_settings.database_path,item["id"],{"stable_key":"user.issue.second","title":"问题二","description":"描述二","corrective_suggestion":"建议二","default_severity":"轻微"}))
    copied = run(copy_issue(test_settings.database_path,second["id"],"user.issue.copy"))
    assert copied["id"] != second["id"] and copied["stable_key"] == "user.issue.copy"
    fourth = run(create_issue(test_settings.database_path,item["id"],{"stable_key":"user.issue.fourth","title":"问题四","description":"描述四","corrective_suggestion":"建议四","default_severity":"一般"}))
    with pytest.raises(TemplateServiceError) as maximum:
        run(create_issue(test_settings.database_path,item["id"],{"stable_key":"user.issue.fifth","title":"问题五","description":"描述五","corrective_suggestion":"建议五","default_severity":"一般"}))
    assert maximum.value.code == "ISSUE_TEMPLATE_ENABLED_LIMIT"
    enabled = run(update_template_entity(test_settings.database_path,"items",item["id"],{"enabled":True}))
    assert enabled["enabled"] == 1
    for issue in (first, second, copied, fourth):
        run(update_issue(test_settings.database_path,issue["id"],{"enabled":False}))
    refreshed = next(value for value in run(checklist_library(test_settings.database_path)) if value["id"] == checklist["id"])["groups"][0]["items"][0]
    assert refreshed["coverage_diagnostic"]["code"] == "ISSUE_TEMPLATE_COVERAGE_EMPTY"


def test_template_service_enforces_starred_severity_without_ui(test_settings) -> None:
    run(initialize_database(test_settings.database_path,test_settings.migrations_dir,test_settings.bootstrap_seed_path,test_settings.seed_schema_path))
    checklist = run(create_template_entity(test_settings.database_path,"checklists",None,{"stable_key":"user.star-checklist","name":"星标规则检查单","audit_kind":"product","display_order":99}))
    group = run(create_template_entity(test_settings.database_path,"groups",checklist["id"],{"stable_key":"user.star-group","name":"星标分组","display_order":1}))
    item = run(create_template_entity(test_settings.database_path,"items",group["id"],{"stable_key":"user.star-item","display_no":"1","content":"星标检查项","starred":True,"display_order":1}))
    with pytest.raises(TemplateServiceError) as invalid:
        run(create_issue(test_settings.database_path,item["id"],{"stable_key":"user.star-issue-invalid","title":"错误严重度","description":"描述","corrective_suggestion":"建议","default_severity":"一般"}))
    assert invalid.value.code == "ISSUE_SEVERITY_STARRED_REQUIRED"
    severe = run(create_issue(test_settings.database_path,item["id"],{"stable_key":"user.star-issue","title":"严重问题","description":"描述","corrective_suggestion":"建议","default_severity":"严重"}))
    with pytest.raises(TemplateServiceError) as changed:
        run(update_issue(test_settings.database_path,severe["id"],{"default_severity":"轻微"}))
    assert changed.value.code == "ISSUE_SEVERITY_STARRED_REQUIRED"
    with pytest.raises(TemplateServiceError) as immutable:
        run(update_template_entity(test_settings.database_path,"items",item["id"],{"starred":False}))
    assert immutable.value.code == "TEMPLATE_ITEM_STARRED_IMMUTABLE"


def test_import_conflict_repeat_stale_and_partial_failure_are_atomic(test_settings, tmp_path: Path) -> None:
    run(initialize_database(test_settings.database_path,test_settings.migrations_dir,test_settings.bootstrap_seed_path,test_settings.seed_schema_path))
    rows = list(run(flat_template_rows(test_settings.database_path)))
    changed = dict(rows[0]); changed["title"] = "事务修改"
    path = build_template_library(tuple([changed, *rows[1:]]), tmp_path / "transaction.xlsx")
    preview = run(create_import_preview(test_settings.database_path,path.read_bytes()))
    run(update_issue(test_settings.database_path,rows[1]["issue_id"],{"title":"预览后变化"}))
    with pytest.raises(TemplateServiceError) as stale:
        run(commit_import_preview(test_settings.database_path,preview["preview_token"]))
    assert stale.value.code == "TEMPLATE_IMPORT_STALE_PREVIEW"
    assert run(flat_template_rows(test_settings.database_path))[0]["title"] != "事务修改"

    current = list(run(flat_template_rows(test_settings.database_path)))
    conflict = dict(current[0]); conflict["item_key"] = current[-1]["item_key"]
    conflict["group_key"] = current[-1]["group_key"]
    conflict["checklist_key"] = current[-1]["checklist_key"]
    conflict["checklist_name"] = current[-1]["checklist_name"]
    conflict["checklist_version"] = current[-1]["checklist_version"]
    conflict["audit_kind"] = current[-1]["audit_kind"]
    conflict["checklist_enabled"] = current[-1]["checklist_enabled"]
    conflict["checklist_order"] = current[-1]["checklist_order"]
    conflict["group_name"] = current[-1]["group_name"]
    conflict["group_version"] = current[-1]["group_version"]
    conflict["group_enabled"] = current[-1]["group_enabled"]
    conflict["group_order"] = current[-1]["group_order"]
    conflict["display_no"] = current[-1]["display_no"]
    conflict["content"] = current[-1]["content"]
    conflict["starred"] = current[-1]["starred"]
    conflict["item_version"] = current[-1]["item_version"]
    conflict["item_enabled"] = current[-1]["item_enabled"]
    conflict["item_order"] = current[-1]["item_order"]
    conflict_path = build_template_library(tuple([conflict, *current[1:]]),tmp_path/"conflict.xlsx")
    conflict_preview = run(create_import_preview(test_settings.database_path,conflict_path.read_bytes()))
    assert any(change["kind"] == "conflict" for change in conflict_preview["changes"])
    before = run(flat_template_rows(test_settings.database_path))
    with pytest.raises(TemplateServiceError) as conflict_error:
        run(commit_import_preview(test_settings.database_path,conflict_preview["preview_token"]))
    assert conflict_error.value.code == "TEMPLATE_IMPORT_CONFLICT"
    assert run(flat_template_rows(test_settings.database_path)) == before

    good = dict(current[0]); good["title"] = "一次提交"
    good_path = build_template_library(tuple([good,*current[1:]]),tmp_path/"good.xlsx")
    good_preview = run(create_import_preview(test_settings.database_path,good_path.read_bytes()))
    run(commit_import_preview(test_settings.database_path,good_preview["preview_token"]))
    with pytest.raises(TemplateServiceError) as repeated:
        run(commit_import_preview(test_settings.database_path,good_preview["preview_token"]))
    assert repeated.value.code == "TEMPLATE_IMPORT_PREVIEW_INVALID"


def test_bulk_import_uses_fixed_select_batches(test_settings, tmp_path: Path, monkeypatch) -> None:
    run(initialize_database(test_settings.database_path,test_settings.migrations_dir,test_settings.bootstrap_seed_path,test_settings.seed_schema_path))
    rows = list(run(flat_template_rows(test_settings.database_path)))
    rows[-1] = {**rows[-1], "title": "批量尾行修改"}
    path = build_template_library(tuple(rows),tmp_path/"bulk.xlsx")
    preview = run(create_import_preview(test_settings.database_path,path.read_bytes()))
    original_connect = template_service.connect
    statements: list[str] = []

    @asynccontextmanager
    async def counted_connect(database_path):
        async with original_connect(database_path) as connection:
            await connection.set_trace_callback(statements.append)
            yield connection

    monkeypatch.setattr(template_service,"connect",counted_connect)
    started = time.perf_counter()
    result = run(commit_import_preview(test_settings.database_path,preview["preview_token"]))
    elapsed = time.perf_counter() - started
    selects = [sql for sql in statements if sql.lstrip().upper().startswith("SELECT")]
    assert result["applied_changes"] == 1
    assert len(rows) >= 900 and len(selects) <= 12 and elapsed < 5


def test_template_api_crud_restore_and_export(test_settings) -> None:
    run(initialize_database(test_settings.database_path,test_settings.migrations_dir,test_settings.bootstrap_seed_path,test_settings.seed_schema_path))
    app = create_app(test_settings)

    async def exercise() -> None:
        transport = httpx.ASGITransport(app=app)
        async with httpx.AsyncClient(transport=transport,base_url="http://test") as client:
            listed = await client.get("/api/templates/checklists")
            assert listed.status_code == 200 and listed.json()[0]["groups"]
            item = listed.json()[0]["groups"][0]["items"][0]
            response = await client.patch(f"/api/templates/items/{item['id']}",json={"content":"API 编辑内容","display_order":88})
            assert response.status_code == 200 and response.json()["version"].endswith("+user.1")
            restored = await client.post(f"/api/templates/items/{item['id']}/restore")
            assert restored.status_code == 200 and restored.json()["user_modified"] == 0
            exported = await client.post("/api/templates/export")
            assert exported.status_code == 200 and exported.content[:4] == b"PK\x03\x04"

    run(exercise())
