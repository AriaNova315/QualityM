from __future__ import annotations

import copy
import json
import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.importers.project_workbook import (
    WorkbookStructureError,
    normalize_wbs,
    parse_project_workbook,
)
from scripts.generate_project_preview import canonical_json, preview_html


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "项目相关信息表.xlsx"
GOLDEN_JSON = ROOT / "tests" / "fixtures" / "project-preview.golden.json"
GOLDEN_HTML = ROOT / "tests" / "fixtures" / "project-preview.golden.html"


def copy_workbook(tmp_path: Path) -> Path:
    target = tmp_path / "project.xlsx"
    shutil.copyfile(SOURCE, target)
    return target


def diagnostic_codes(payload: dict) -> set[str]:
    return {diagnostic["code"] for diagnostic in payload["diagnostics"]}


def test_current_project_sample_matches_golden_baseline() -> None:
    payload = parse_project_workbook(SOURCE)

    assert payload["project"] == {
        "project_name": "相控阵数字样机软件",
        "project_code": "03046",
        "project_leader": "赵要浦",
        "qa_person": "高蕾",
        "source": {"sheet": "项目信息", "row": 2},
    }
    assert payload["schedule_row_count"] == 200
    assert [(phase["wbs"], phase["name"]) for phase in payload["phases"]] == [
        ("1", "系统分析与设计"),
        ("2", "需求分析阶段"),
        ("3", "软件设计阶段"),
        ("4", "软件编码与单元测试阶段"),
        ("6", "软件合格性测试及集成试运行阶段"),
        ("7", "验收与交付"),
    ]
    assert len(payload["tracking_nodes"]) == 7
    assert len(payload["quality_activities"]) == 56
    assert len(payload["config_items"]) == 26
    assert payload["diagnostics"] == []
    assert payload["excluded_top_level"][0]["name"] == "公司质量验证"
    assert payload["excluded_top_level"][0]["reason"] == "excluded_company_validation"


def test_formula_values_and_source_formulas_are_both_preserved() -> None:
    payload = parse_project_workbook(SOURCE)
    activity = next(item for item in payload["quality_activities"] if item["wbs"] == "2.6.9")
    first_item = payload["config_items"][0]

    assert activity["linked_items"] == (
        "相控阵数字样机软件开发计划",
        "相控阵数字样机软件测量分析计划",
        "相控阵数字样机软件质量保证计划",
        "相控阵数字样机软件配置管理计划",
    )
    assert activity["linked_items_formula"] == "=E45"
    assert first_item["phase_wbs"] == "1"
    assert first_item["phase_wbs_formula"] == "=项目进度计划!A2"


def test_only_direct_children_of_quality_parent_are_selected(tmp_path: Path) -> None:
    path = copy_workbook(tmp_path)
    workbook = load_workbook(path)
    sheet = workbook["项目进度计划"]
    sheet.append(["1.5.1.1", "产品审核-孙级活动", sheet["C18"].value, sheet["D18"].value, None, 1, 1, "高蕾"])
    workbook.save(path)

    payload = parse_project_workbook(path)
    assert all(activity["name"] != "产品审核-孙级活动" for activity in payload["quality_activities"])
    assert len(payload["quality_activities"]) == 56


def test_duplicate_item_and_multiple_qa_resources_have_stable_diagnostics(tmp_path: Path) -> None:
    path = copy_workbook(tmp_path)
    workbook = load_workbook(path)
    config = workbook["配置项清单"]
    config.append([27, config["B2"].value, "DUP", config["D2"].value, "文档"])
    schedule = workbook["项目进度计划"]
    schedule["H18"] = "高蕾,赵要浦"
    workbook.save(path)

    payload = parse_project_workbook(path)
    assert diagnostic_codes(payload) >= {
        "IMPORT_DUPLICATE_CONFIG_ITEM",
        "IMPORT_MULTIPLE_QA_RESOURCES",
    }
    duplicate = next(d for d in payload["diagnostics"] if d["code"] == "IMPORT_DUPLICATE_CONFIG_ITEM")
    multiple = next(d for d in payload["diagnostics"] if d["code"] == "IMPORT_MULTIPLE_QA_RESOURCES")
    assert duplicate["location"] == {"sheet": "配置项清单", "row": 28}
    assert multiple["location"] == {"sheet": "项目进度计划", "row": 18}


def test_unknown_activity_and_product_audit_without_link_are_preview_blockers(
    tmp_path: Path,
) -> None:
    path = copy_workbook(tmp_path)
    workbook = load_workbook(path)
    schedule = workbook["项目进度计划"]
    schedule["B18"] = "审核软件研制任务书"
    schedule["E61"] = None
    workbook.save(path)

    payload = parse_project_workbook(path)
    unknown = next(
        diagnostic
        for diagnostic in payload["diagnostics"]
        if diagnostic["code"] == "IMPORT_UNKNOWN_QUALITY_ACTIVITY"
        and diagnostic["location"]["row"] == 18
    )
    missing_link = next(
        diagnostic
        for diagnostic in payload["diagnostics"]
        if diagnostic["code"] == "IMPORT_PRODUCT_AUDIT_WITHOUT_CONFIG_ITEM"
        and diagnostic["location"]["row"] == 61
    )
    assert unknown["severity"] == "blocker"
    assert missing_link["severity"] == "blocker"
    assert unknown["location"] == {"sheet": "项目进度计划", "row": 18}
    assert missing_link["location"] == {"sheet": "项目进度计划", "row": 61}


def test_missing_sheet_and_column_are_blockers(tmp_path: Path) -> None:
    path = copy_workbook(tmp_path)
    workbook = load_workbook(path)
    del workbook["配置项清单"]
    workbook["项目进度计划"]["D1"] = "错误列名"
    workbook.save(path)

    with pytest.raises(WorkbookStructureError) as error:
        parse_project_workbook(path)
    assert {diagnostic.code for diagnostic in error.value.diagnostics} == {
        "IMPORT_MISSING_SHEET",
        "IMPORT_MISSING_COLUMN",
    }


def test_missing_date_is_a_blocker(tmp_path: Path) -> None:
    path = copy_workbook(tmp_path)
    workbook = load_workbook(path)
    workbook["项目进度计划"]["D18"] = None
    workbook.save(path)

    payload = parse_project_workbook(path)
    diagnostic = next(d for d in payload["diagnostics"] if d["code"] == "IMPORT_MISSING_SCHEDULE_DATE")
    assert diagnostic["severity"] == "blocker"
    assert diagnostic["location"] == {"sheet": "项目进度计划", "row": 18}


@pytest.mark.parametrize(
    "value, expected",
    [(1, "1"), (1.1, "1.1"), (1.01, "1.01"), (" 1.10 ", "1.10"), ("2 . 6 . 10", "2.6.10")],
)
def test_wbs_normalization_does_not_introduce_float_artifacts(value, expected: str) -> None:
    assert normalize_wbs(value) == expected


def test_golden_payload_is_json_serializable_and_source_located() -> None:
    payload = parse_project_workbook(SOURCE)
    serialized = json.dumps(payload, ensure_ascii=False, sort_keys=True)
    assert "项目进度计划" in serialized
    assert all("source" in phase and phase["source"]["row"] >= 2 for phase in payload["phases"])
    assert all("source" in item and item["source"]["row"] >= 2 for item in payload["config_items"])


def test_same_config_item_name_with_different_version_is_allowed() -> None:
    payload = parse_project_workbook(SOURCE)
    source_code = [item for item in payload["config_items"] if item["name"].endswith("源代码")]
    assert len(source_code) == 5
    assert len({item["version"] for item in source_code}) == 5
    assert "IMPORT_DUPLICATE_CONFIG_ITEM" not in diagnostic_codes(payload)


def test_golden_preview_artifacts_are_current() -> None:
    payload = parse_project_workbook(SOURCE)
    assert GOLDEN_JSON.read_text(encoding="utf-8") == canonical_json(payload)
    assert GOLDEN_HTML.read_text(encoding="utf-8") == preview_html(payload)
