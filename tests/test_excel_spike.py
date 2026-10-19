from __future__ import annotations

import hashlib
import shutil
from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook

from app.excel.external_links import add_external_workbook_link
from app.excel.spike_builder import (
    AUDIT_REPORT_SHEET,
    CHECKLIST_SHEET,
    PHASE_REPORT_SHEET,
    TREND_SHEET,
    ExcelSpikeSnapshot,
    build_excel_linkage_spike,
)
from app.excel.validator import validate_excel_linkage_spike


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE_DIR = ROOT / "temp"


@pytest.fixture
def spike_artifacts(tmp_path: Path):
    return build_excel_linkage_spike(
        ExcelSpikeSnapshot(),
        output_dir=tmp_path / "step5",
        template_dir=TEMPLATE_DIR,
    )


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def test_spike_preserves_authoritative_templates_and_writes_only_output(tmp_path: Path) -> None:
    sources = tuple(TEMPLATE_DIR.glob("*.xlsx"))
    before = {path.name: _sha256(path) for path in sources}
    output_dir = tmp_path / "isolated"

    build_excel_linkage_spike(
        ExcelSpikeSnapshot(),
        output_dir=output_dir,
        template_dir=TEMPLATE_DIR,
    )

    assert {path.name: _sha256(path) for path in sources} == before
    assert {path.suffix for path in output_dir.iterdir()} == {".xlsx"}


def test_phase_workbook_reopens_with_formula_chain_style_and_editable_sheets(spike_artifacts) -> None:
    workbook = load_workbook(spike_artifacts.phase_workbook, data_only=False)
    assert workbook.sheetnames == [CHECKLIST_SHEET, AUDIT_REPORT_SHEET, PHASE_REPORT_SHEET]
    assert workbook.calculation.calcMode == "auto"
    assert workbook.calculation.fullCalcOnLoad is True
    assert workbook.calculation.forceFullCalc is True

    checklist = workbook[CHECKLIST_SHEET]
    report = workbook[AUDIT_REPORT_SHEET]
    phase = workbook[PHASE_REPORT_SHEET]
    assert checklist["A1"].value == "软件需求规格说明检查单"
    assert checklist["A1"].style_id != 0
    assert checklist.protection.sheet is False
    assert report.protection.sheet is False
    assert phase.protection.sheet is False
    assert f"'{CHECKLIST_SHEET}'!$H$7" in report["C13"].value
    assert report["D10"].value == '=COUNTIF(C13:C16,"?*")'
    assert f"'{AUDIT_REPORT_SHEET}'!$D$10" in phase["A91"].value
    assert phase["L101"].value == "=SUM(C101:K101)"
    assert phase.print_area == f"'{PHASE_REPORT_SHEET}'!$A$1:$L$104"
    workbook.close()


def test_trend_workbook_contains_relative_external_link_and_chart(spike_artifacts) -> None:
    validation = validate_excel_linkage_spike(
        spike_artifacts.phase_workbook,
        spike_artifacts.trend_workbook,
    )
    assert validation.external_link_parts == ("xl/externalLinks/externalLink1.xml",)
    assert validation.chart_parts == ("xl/charts/chart1.xml",)

    workbook = load_workbook(spike_artifacts.trend_workbook, data_only=False, keep_links=True)
    worksheet = workbook[TREND_SHEET]
    assert worksheet["B4"].value == "='[1]阶段质量审核工作报告'!$L$101"
    assert len(worksheet._charts) == 1
    link = workbook._external_links[0]
    assert link.file_link.Target == spike_artifacts.phase_workbook.name
    assert link.file_link.TargetMode == "External"
    workbook.close()

    with ZipFile(spike_artifacts.trend_workbook) as archive:
        relationships = archive.read("xl/externalLinks/_rels/externalLink1.xml.rels").decode()
        workbook_relationships = archive.read("xl/_rels/workbook.xml.rels").decode()
        content_types = archive.read("[Content_Types].xml").decode()
        assert f'Target="{spike_artifacts.phase_workbook.name}"' in relationships
        assert "TargetMode=\"External\"" in relationships
        assert "relationships/externalLink" in workbook_relationships
        assert "spreadsheetml.externalLink+xml" in content_types
        assert "xl/drawings/drawing1.xml" in archive.namelist()


def test_moving_whole_output_directory_preserves_relative_link(spike_artifacts, tmp_path: Path) -> None:
    moved_dir = tmp_path / "moved" / "complete-export"
    shutil.copytree(spike_artifacts.phase_workbook.parent, moved_dir)
    moved_phase = moved_dir / spike_artifacts.phase_workbook.name
    moved_trend = moved_dir / spike_artifacts.trend_workbook.name

    workbook = load_workbook(moved_trend, data_only=False, keep_links=True)
    assert moved_phase.is_file()
    assert workbook._external_links[0].file_link.Target == moved_phase.name
    assert workbook[TREND_SHEET]["B4"].value == "='[1]阶段质量审核工作报告'!$L$101"
    workbook.close()


def test_excel_smoke_script_refuses_ambiguous_live_workbooks() -> None:
    source = (ROOT / "scripts" / "excel_smoke.applescript").read_text(encoding="utf-8")
    assert "requires no pre-existing workbooks" in source
    assert "active workbook" not in source
    assert "close every workbook" not in source


@pytest.mark.parametrize("target", ("", "/tmp/phase.xlsx", "../phase.xlsx", "nested/phase.xlsx", "phase.xls"))
def test_external_link_adapter_rejects_unsafe_targets(target: str) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    with pytest.raises(ValueError):
        add_external_workbook_link(
            workbook,
            target=target,
            sheet_names=(PHASE_REPORT_SHEET,),
        )
