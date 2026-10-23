from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook

from app.excel import formula_cache


def _workbook(path: Path) -> None:
    workbook = Workbook()
    workbook.active.title = "统计"
    workbook.active["A1"] = "=1+2"
    workbook.active["A2"] = '="结果"'
    workbook.save(path)
    workbook.close()


def test_caches_replace_only_after_all_zip_handles_close(tmp_path, monkeypatch):
    path = tmp_path / "质量记录.xlsx"
    _workbook(path)
    opened = set()
    replace = formula_cache.os.replace

    class TrackedZipFile(ZipFile):
        def __enter__(self):
            opened.add(self)
            return super().__enter__()

        def __exit__(self, *args):
            try:
                return super().__exit__(*args)
            finally:
                opened.remove(self)

    def windows_replace(source, target):
        # Reproduce Windows' refusal to replace a file with an open ZIP handle.
        assert not opened, "Workbook ZIP handles must close before replacement"
        return replace(source, target)

    monkeypatch.setattr(formula_cache, "ZipFile", TrackedZipFile)
    monkeypatch.setattr(formula_cache.os, "replace", windows_replace)
    formula_cache.write_formula_caches(path, {"统计": {"A1": 3, "A2": "结果"}})

    values = load_workbook(path, data_only=True)
    formulas = load_workbook(path, data_only=False)
    try:
        assert values["统计"]["A1"].value == 3
        assert values["统计"]["A2"].value == "结果"
        assert formulas["统计"]["A1"].value == "=1+2"
    finally:
        values.close()
        formulas.close()
    assert set(tmp_path.iterdir()) == {path}


@pytest.mark.parametrize("failure", ("missing_cell", "replace_failure"))
def test_cache_failure_preserves_source_and_cleans_temporary_file(tmp_path, monkeypatch, failure):
    path = tmp_path / "质量记录.xlsx"
    _workbook(path)
    original = path.read_bytes()
    if failure == "replace_failure":
        def fail_replace(*args):
            raise PermissionError("workbook is locked")
        monkeypatch.setattr(formula_cache.os, "replace", fail_replace)
        error, cell = PermissionError, "A1"
    else:
        error, cell = ValueError, "Z99"
    with pytest.raises(error):
        formula_cache.write_formula_caches(path, {"统计": {cell: 3}})
    assert path.read_bytes() == original
    assert set(tmp_path.iterdir()) == {path}
