from __future__ import annotations

import asyncio
import sqlite3
from dataclasses import FrozenInstanceError, replace
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.application.export_snapshot_service import (
    ExportSnapshotError,
    load_export_snapshot,
)
from app.application.audit_service import set_item_conclusion
from app.application.import_preview_service import create_preview
from app.application.project_import_service import commit_preview
from app.config import Settings
from app.database import initialize_database
from app.excel.phase_workbook_builder import (
    CALCULATION_SHEET,
    DIRECTORY_SHEET,
    PHASE_REPORT_SHEET,
    build_all_phase_workbooks,
    build_phase_workbook,
)
from app.excel.validator import validate_phase_workbook
from app.excel.sheet_names import safe_sheet_name


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "项目相关信息表.xlsx"


def run(coroutine):
    return asyncio.run(coroutine)


@pytest.fixture(scope="module")
def step8_context(tmp_path_factory):
    project_root = tmp_path_factory.mktemp("step8") / "qualitym-test"
    settings = Settings(
        app_name="QualityM",
        app_version="0.1.0",
        host="127.0.0.1",
        port=8003,
        reload=False,
        project_root=project_root,
        database_path=project_root / "qualitym.db",
        generated_dir=project_root / "generated",
        static_dir=ROOT / "app" / "static",
        migrations_dir=ROOT / "migrations",
        bootstrap_seed_path=ROOT / "seeds" / "checklist-library.json",
        seed_schema_path=ROOT / "seeds" / "schema" / "checklist-library.schema.json",
        import_temp_dir=project_root / "var" / "imports",
        upload_max_bytes=25 * 1024 * 1024,
        import_preview_ttl_seconds=3600,
    )
    run(
        initialize_database(
            settings.database_path,
            settings.migrations_dir,
            settings.bootstrap_seed_path,
            settings.seed_schema_path,
        )
    )
    preview = run(
        create_preview(
            settings.database_path,
            settings.import_temp_dir,
            SOURCE.name,
            SOURCE.read_bytes(),
            settings.import_preview_ttl_seconds,
            settings.upload_max_bytes,
        )
    )
    committed = run(
        commit_preview(settings.database_path, preview.token, {}, {}, False, "0.0")
    )
    with sqlite3.connect(settings.database_path) as connection:
        audit_id, revision, item_id = connection.execute(
            """
            SELECT ai.id, ai.revision, cii.id
            FROM audit_instances ai
            JOIN checklist_instance_groups cig ON cig.audit_instance_id=ai.id
            JOIN checklist_instance_items cii ON cii.instance_group_id=cig.id
            WHERE ai.project_version_id=? AND ai.report_scope='phase'
              AND cii.starred_snapshot=0
              AND (SELECT COUNT(*) FROM checklist_instance_issue_templates ciit
                   WHERE ciit.checklist_instance_item_id=cii.id) >= 2
            ORDER BY ai.sequence_no, cii.id LIMIT 1
            """,
            (committed.project_version_id,),
        ).fetchone()
        candidate_ids = [
            row[0]
            for row in connection.execute(
                "SELECT id FROM checklist_instance_issue_templates "
                "WHERE checklist_instance_item_id=? ORDER BY display_order,id LIMIT 2",
                (item_id,),
            ).fetchall()
        ]
    run(
        set_item_conclusion(
            settings.database_path,
            int(item_id),
            int(revision),
            "不符合",
            [int(value) for value in candidate_ids],
            False,
        )
    )
    snapshot = run(load_export_snapshot(settings.database_path, committed.project_version_id))
    artifacts = build_all_phase_workbooks(snapshot, project_root / "step8-output")
    return settings, committed, snapshot, artifacts


def test_export_snapshot_is_complete_immutable_and_deterministic(step8_context) -> None:
    settings, committed, snapshot, _artifacts = step8_context
    assert len(snapshot.phases) == 6
    assert sum(len(phase.audits) for phase in snapshot.phases) == 52
    assert sum(
        audit.report_scope == "company"
        for phase in snapshot.phases
        for audit in phase.audits
    ) == 1
    assert len(snapshot.content_sha256) == 64
    assert run(load_export_snapshot(settings.database_path, committed.project_version_id)) == snapshot
    with pytest.raises(FrozenInstanceError):
        snapshot.project_name = "changed"
    with sqlite3.connect(settings.database_path) as connection:
        assert connection.execute("PRAGMA journal_mode").fetchone()[0] == "wal"


def test_export_snapshot_reports_missing_project_version(step8_context) -> None:
    settings, _committed, _snapshot, _artifacts = step8_context
    with pytest.raises(ExportSnapshotError) as error:
        run(load_export_snapshot(settings.database_path, 999999))
    assert error.value.code == "PROJECT_VERSION_NOT_FOUND"


def test_current_sample_builds_every_phase_and_reopens_cleanly(step8_context) -> None:
    _settings, _committed, snapshot, artifacts = step8_context
    assert len(artifacts) == len(snapshot.phases) == 6
    for artifact, phase in zip(artifacts, snapshot.phases, strict=True):
        validation = validate_phase_workbook(artifact.path)
        expected_audits = sum(audit.report_scope == "phase" for audit in phase.audits)
        expected_issues = sum(
            len(item.issues)
            for audit in phase.audits
            if audit.report_scope == "phase"
            for group in audit.groups
            for item in group.items
        )
        assert validation.audit_count == expected_audits == artifact.audit_count
        assert validation.issue_count == expected_issues == artifact.issue_count
        assert validation.formula_count >= expected_issues * 8
        assert artifact.path.name.endswith("阶段质量审核记录.xlsx")


def test_workbook_has_formula_chain_hyperlinks_editable_sheets_and_print_layout(step8_context) -> None:
    _settings, _committed, _snapshot, artifacts = step8_context
    artifact = next(item for item in artifacts if item.audit_count > 0 and item.issue_count > 0)
    workbook = load_workbook(artifact.path, data_only=False)
    try:
        assert workbook.sheetnames[:2] == [DIRECTORY_SHEET, PHASE_REPORT_SHEET]
        assert workbook[CALCULATION_SHEET].sheet_state == "hidden"
        checklist_name = next(name for name in workbook.sheetnames if name.startswith("C"))
        report_name = next(name for name in workbook.sheetnames if name.startswith("R"))
        checklist = workbook[checklist_name]
        report = workbook[report_name]
        phase = workbook[PHASE_REPORT_SHEET]
        assert not checklist.protection.sheet
        assert not report.protection.sheet
        assert not phase.protection.sheet
        assert checklist.print_area and report.print_area and phase.print_area
        assert workbook[DIRECTORY_SHEET]["C5"].hyperlink.target.startswith("#'")
        assert report["B1"].value == "过程/产品审核报告"
        assert report.page_setup.orientation == "portrait"
        assert report.row_dimensions[8].height >= 128.5
        assert report.row_dimensions[13].height >= 72
        report_last_row = max(
            cell.row
            for row in report.iter_rows(min_col=2, max_col=11)
            for cell in row
            if cell.value is not None
        )
        assert all(report.cell(row, 2).border.left.style == "thick" for row in range(3, report_last_row + 1))
        assert all(report.cell(row, 11).border.right.style == "thick" for row in range(3, report_last_row + 1))
        assert all(report.cell(3, column).border.top.style == "thick" for column in range(2, 12))
        assert all(report.cell(report_last_row, column).border.bottom.style == "thick" for column in range(2, 12))
        assert checklist["A1"].font.name in {"Microsoft YaHei", "黑体", "宋体"}
        assert "不符合项问题明细" not in {
            cell.value for row in checklist.iter_rows() for cell in row
        }
        assert any(
            cell.data_type == "f" and CALCULATION_SHEET in str(cell.value)
            for row in phase.iter_rows()
            for cell in row
        )
        assert any(
            isinstance(cell.value, str) and "待确认的建议值" in cell.value
            for row in report.iter_rows()
            for cell in row
        )
        assert "产品NC统计" in {
            cell.value for row in phase.iter_rows() for cell in row
        }
        assert "过程NC统计" in {
            cell.value for row in phase.iter_rows() for cell in row
        }
    finally:
        workbook.close()


def test_phase_summary_keeps_all_ten_trend_categories(step8_context) -> None:
    _settings, _committed, _snapshot, artifacts = step8_context
    workbook = load_workbook(artifacts[0].path, data_only=False)
    try:
        calculation = workbook[CALCULATION_SHEET]
        values = [cell.value for row in calculation.iter_rows() for cell in row]
        assert {"RDM", "PP", "PMC", "CM", "MT", "MPM", "VV", "LD", "II", "工作产品"}.issubset(values)
    finally:
        workbook.close()


def test_zero_audit_phase_still_has_valid_directory_node_blocks_and_statistics(
    step8_context, tmp_path: Path
) -> None:
    _settings, _committed, snapshot, _artifacts = step8_context
    empty_phase = replace(snapshot.phases[0], audits=())
    empty_snapshot = replace(snapshot, phases=(empty_phase,))
    path = tmp_path / "empty-phase.xlsx"
    artifact = build_phase_workbook(empty_snapshot, empty_phase, path)
    validation = validate_phase_workbook(path)
    assert (artifact.audit_count, artifact.issue_count) == (0, 0)
    assert (validation.audit_count, validation.issue_count) == (0, 0)
    workbook = load_workbook(path, data_only=False)
    try:
        phase_report = workbook[PHASE_REPORT_SHEET]
        values = [cell.value for row in phase_report.iter_rows() for cell in row]
        assert "本阶段无审核活动。" in values
        assert "阶段汇总统计：" in values
    finally:
        workbook.close()


def test_snapshot_isolated_from_later_database_edits(step8_context) -> None:
    settings, committed, snapshot, _artifacts = step8_context
    original_hash = snapshot.content_sha256
    with sqlite3.connect(settings.database_path) as connection:
        connection.execute(
            "UPDATE audit_instances SET auditor='快照之后修改', revision=revision+1 "
            "WHERE project_version_id=? AND sequence_no=1",
            (committed.project_version_id,),
        )
        connection.commit()
    assert snapshot.content_sha256 == original_hash
    refreshed = run(load_export_snapshot(settings.database_path, committed.project_version_id))
    assert refreshed.content_sha256 != original_hash
    assert refreshed.phases[0].audits[0].auditor == "快照之后修改"


def test_split_issues_on_one_checklist_item_are_all_exported(
    step8_context, tmp_path: Path
) -> None:
    _settings, _committed, snapshot, _artifacts = step8_context
    source_phase = next(
        phase
        for phase in snapshot.phases
        if any(item.issues for audit in phase.audits for group in audit.groups for item in group.items)
    )
    source_audit = next(
        audit
        for audit in source_phase.audits
        if any(item.issues for group in audit.groups for item in group.items)
    )
    source_group = next(group for group in source_audit.groups if any(item.issues for item in group.items))
    source_item = next(item for item in source_group.items if item.issues)
    duplicate = replace(
        source_item.issues[0],
        id=9_999_999,
        title="拆分问题",
        description="同一检查项拆分出的第二条问题。",
    )
    changed_item = replace(source_item, issues=source_item.issues + (duplicate,))
    changed_group = replace(
        source_group,
        items=tuple(changed_item if item.id == source_item.id else item for item in source_group.items),
    )
    changed_audit = replace(
        source_audit,
        groups=tuple(changed_group if group.id == source_group.id else group for group in source_audit.groups),
        tracking_node_type=None,
        tracking_node_date=None,
    )
    changed_phase = replace(source_phase, audits=(changed_audit,))
    changed_snapshot = replace(snapshot, phases=(changed_phase,))
    path = tmp_path / "split-issues.xlsx"
    artifact = build_phase_workbook(changed_snapshot, changed_phase, path)
    validation = validate_phase_workbook(path)
    expected_issues = sum(
        len(item.issues) for group in changed_audit.groups for item in group.items
    )
    assert artifact.issue_count == validation.issue_count == expected_issues
    workbook = load_workbook(path, data_only=False)
    try:
        calculation_values = [
            cell.value for row in workbook[CALCULATION_SHEET].iter_rows() for cell in row
        ]
        assert any("未匹配节点" == value for value in calculation_values)
        report = workbook[next(name for name in workbook.sheetnames if name.startswith("R"))]
        assert report["C15"].value == "同一检查项拆分出的第二条问题。"
        assert report.row_dimensions[15].height >= 72
    finally:
        workbook.close()


def test_business_text_cannot_become_formula_and_sheet_names_are_safe(
    step8_context, tmp_path: Path
) -> None:
    _settings, _committed, snapshot, _artifacts = step8_context
    phase = replace(snapshot.phases[0], audits=())
    changed = replace(snapshot, project_name="=HYPERLINK(\"unsafe\")", phases=(phase,))
    path = tmp_path / "literal-boundary.xlsx"
    build_phase_workbook(changed, phase, path)
    workbook = load_workbook(path, data_only=False)
    try:
        assert workbook[DIRECTORY_SHEET]["A2"].data_type != "f"
        assert workbook[PHASE_REPORT_SHEET]["B3"].value.startswith("'")
    finally:
        workbook.close()
    used: set[str] = set()
    first = safe_sheet_name("[超长/非法:*?]" * 10, prefix="R001-", used=used)
    second = safe_sheet_name("[超长/非法:*?]" * 10, prefix="R001-", used=used)
    assert len(first) <= 31 and len(second) <= 31
    assert first != second
    assert not any(character in first + second for character in "\\/*?:[]")


def test_same_date_same_type_tracking_nodes_do_not_duplicate_audits(
    step8_context, tmp_path: Path
) -> None:
    _settings, _committed, snapshot, _artifacts = step8_context
    source_phase = next(
        phase
        for phase in snapshot.phases
        if phase.tracking_nodes and any(audit.tracking_node_id for audit in phase.audits)
    )
    source_node = source_phase.tracking_nodes[0]
    duplicate_node = replace(source_node, id=9_999_998, source_wbs="duplicate-node")
    changed_phase = replace(
        source_phase,
        tracking_nodes=source_phase.tracking_nodes + (duplicate_node,),
    )
    changed_snapshot = replace(snapshot, phases=(changed_phase,))
    path = tmp_path / "duplicate-node-key.xlsx"
    build_phase_workbook(changed_snapshot, changed_phase, path)
    workbook = load_workbook(path, data_only=False)
    try:
        values = [
            cell.value
            for row in workbook[PHASE_REPORT_SHEET].iter_rows()
            for cell in row
        ]
        assert any(
            isinstance(value, str) and "无审核活动" in value
            for value in values
        )
    finally:
        workbook.close()
