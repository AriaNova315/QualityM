from __future__ import annotations

from dataclasses import replace
import asyncio
from pathlib import Path
import pytest
from openpyxl import load_workbook

from app.application.export_snapshot_service import load_export_snapshot
from app.application.import_preview_service import create_preview
from app.application.project_import_service import commit_preview
from app.database import initialize_database

from app.excel.company_trend_builder import (
    TREND_SHEET, build_complete_export_package,
)
from app.excel.phase_workbook_builder import build_all_phase_workbooks
from app.excel.validator import validate_company_trend_workbooks


def run(coroutine):
    return asyncio.run(coroutine)


@pytest.fixture
def export_context(test_settings):
    root = Path(__file__).resolve().parents[1]
    run(initialize_database(test_settings.database_path, test_settings.migrations_dir, test_settings.bootstrap_seed_path, test_settings.seed_schema_path))
    source = root / "项目相关信息表.xlsx"
    preview = run(create_preview(test_settings.database_path, test_settings.import_temp_dir, source.name, source.read_bytes(), test_settings.import_preview_ttl_seconds, test_settings.upload_max_bytes))
    committed = run(commit_preview(test_settings.database_path, preview.token, {}, {}, False, "0.0"))
    return run(load_export_snapshot(test_settings.database_path, committed.project_version_id))


def test_complete_package_has_manifest_company_trend_links_and_charts(export_context, tmp_path) -> None:
    snapshot = export_context
    package = build_complete_export_package(snapshot, tmp_path, build_all_phase_workbooks)
    assert package.output_dir.parent == tmp_path
    assert package.company_workbook and package.company_workbook.is_file()
    assert package.trend_workbook.is_file()
    assert set(package.manifest.read_text(encoding="utf-8").splitlines()) == {
        path.name for path in (*package.phase_workbooks, package.company_workbook, package.trend_workbook)
    }
    validate_company_trend_workbooks(package.company_workbook, package.trend_workbook)
    for path in (*package.phase_workbooks, package.company_workbook, package.trend_workbook):
        exported = load_workbook(path, data_only=False, keep_links=True)
        try:
            assert all(not sheet.protection.sheet for sheet in exported.worksheets)
        finally:
            exported.close()
    workbook = load_workbook(package.trend_workbook, data_only=False, keep_links=True)
    try:
        assert len(workbook[TREND_SHEET]._charts) == 2
        assert len(workbook._external_links) == sum(
            any(audit.report_scope == "phase" for audit in phase.audits)
            for phase in snapshot.phases
        )
        assert workbook[TREND_SHEET]["B4"].data_type == "f"
        assert workbook[TREND_SHEET]["Q3"].value == "类别"
        assert workbook[TREND_SHEET]["Q11"].value == "LD"
        assert workbook[TREND_SHEET]["R4"].value == f"=F{4 + len(snapshot.phases)}"
        assert workbook[TREND_SHEET]["S4"].value == f"=IFERROR(R4/$B${4 + len(snapshot.phases)},0)"
        assert workbook[TREND_SHEET]["S4"].number_format == "0.0%"
        line, pie = workbook[TREND_SHEET]._charts
        assert len(line.ser) == 2
        assert len(pie.ser) == 1
    finally:
        workbook.close()


def test_no_company_activity_skips_company_workbook_and_empty_phase_has_no_link(export_context, tmp_path) -> None:
    snapshot = export_context
    phases = tuple(replace(phase, audits=()) if phase.sort_order == 1 else phase for phase in snapshot.phases)
    phases = tuple(replace(phase, audits=tuple(a for a in phase.audits if a.report_scope == "phase")) for phase in phases)
    package = build_complete_export_package(replace(snapshot, phases=phases), tmp_path, build_all_phase_workbooks)
    assert package.company_workbook is None
    workbook = load_workbook(package.trend_workbook, data_only=False, keep_links=True)
    try:
        assert workbook[TREND_SHEET]["B4"].value == 0
        assert len(workbook._external_links) == sum(
            any(audit.report_scope == "phase" for audit in phase.audits)
            for phase in phases
        )
    finally:
        workbook.close()


def test_trend_conclusion_uses_manual_override_and_empty_override_uses_formula(export_context, tmp_path) -> None:
    overridden = replace(export_context, trend_conclusion_override="人工审核结论")
    package = build_complete_export_package(overridden, tmp_path / "override", build_all_phase_workbooks)
    workbook = load_workbook(package.trend_workbook, data_only=False, keep_links=True)
    try:
        conclusion_row = 4 + len(overridden.phases) + 2
        assert workbook[TREND_SHEET].cell(conclusion_row, 1).value == "人工审核结论"
    finally:
        workbook.close()

    package = build_complete_export_package(export_context, tmp_path / "formula", build_all_phase_workbooks)
    workbook = load_workbook(package.trend_workbook, data_only=False, keep_links=True)
    try:
        conclusion_row = 4 + len(export_context.phases) + 2
        assert workbook[TREND_SHEET].cell(conclusion_row, 1).data_type == "f"
    finally:
        workbook.close()


def test_nonzero_snapshot_caches_and_relative_links(export_context, tmp_path):
    from app.application.export_snapshot_service import ExportIssueSnapshot
    from app.excel.phase_workbook_builder import CALCULATION_SHEET
    from zipfile import ZipFile
    import xml.etree.ElementTree as ET
    import shutil

    # Two stages and distinct severities catch partial refresh and all-zero fixtures.
    phases = list(export_context.phases)
    for index, severity in ((0, "严重"), (1, "一般")):
        phase = phases[index]
        audit = next(a for a in phase.audits if a.report_scope == "phase")
        group = audit.groups[0]
        item = group.items[0]
        issue = ExportIssueSnapshot(100 + index, "缓存回归", "缓存回归", severity,
            None, None, None, None, None, None, None, False)
        changed_item = replace(item, conclusion="不符合", issues=(issue,))
        changed_group = replace(group, items=(changed_item, *group.items[1:]))
        changed_audit = replace(audit, groups=(changed_group, *audit.groups[1:]))
        phases[index] = replace(phase, audits=tuple(changed_audit if a.id == audit.id else a for a in phase.audits))
    snapshot = replace(export_context, project_code="QA-02-hyphen", phases=tuple(phases))
    package = build_complete_export_package(snapshot, tmp_path / "original", build_all_phase_workbooks)
    moved = tmp_path / "moved"
    shutil.move(package.output_dir, moved)
    path = moved / package.trend_workbook.name
    validate_company_trend_workbooks(None, path)
    values = load_workbook(path, data_only=True)
    formulas = load_workbook(path, data_only=False)
    try:
        sheet = values[TREND_SHEET]
        assert [sheet[f"B{r}"].value for r in range(4, 10)] == [1, 1, 0, 0, 0, 0]
        assert sheet['B10'].value == 2
        assert sheet['D10'].value == sheet['E10'].value == 1
        assert sheet['A12'].value == '当前项目共发现2项不符合项，其中严重1项。'
        assert sum(sheet[f'S{r}'].value for r in range(4, 14)) == 1
        assert sum(sheet[f'R{r}'].value for r in range(4, 14)) == 2
        for i, link in enumerate(formulas._external_links):
            assert link.file_link.Target == package.phase_workbooks[i].name
            assert (moved / link.file_link.Target).is_file()
            cached = link.externalBook.sheetDataSet.sheetData[0]
            assert cached.row[0].cell[0].v == str(1 if i < 2 else 0)
        assert 'IFERROR' not in formulas[TREND_SHEET]['B4'].value
        assert '$A:$A' not in formulas[TREND_SHEET]['B4'].value
        for i, original in enumerate(package.phase_workbooks):
            book = load_workbook(moved / original.name, data_only=True)
            try:
                summary = {r[0].value: r[1].value for r in book[CALCULATION_SHEET] if isinstance(r[0].value, str)}
                assert summary['NC总数'] == (1 if i < 2 else 0)
            finally:
                book.close()
        with ZipFile(path) as archive:
            ns = {'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart'}
            chart = ET.fromstring(archive.read('xl/charts/chart1.xml'))
            assert [float(v.text) for v in chart.findall('.//c:ser/c:val/c:numRef/c:numCache/c:pt/c:v', ns)] == [1,1,0,0,0,0,1,0,0,0,0,0]
    finally:
        values.close()
        formulas.close()

    # An openpyxl re-save drops caches; the export gate must catch this.
    uncached = load_workbook(path)
    uncached.save(path)
    uncached.close()
    with pytest.raises(ValueError, match="snapshot values are missing"):
        validate_company_trend_workbooks(None, path)
