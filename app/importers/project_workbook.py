from __future__ import annotations

import re
from dataclasses import asdict, dataclass
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.domain.activity_rules import classify_activity
from app.diagnostics import Diagnostic


REQUIRED_HEADERS = {
    "项目信息": ["项目名称", "项目代号", "项目组长", "质量管理员"],
    "项目进度计划": [
        "WBS", "任务名称", "计划开始时间", "计划完成时间", "关联配置项", "资源名称\n(逗号分隔)"
    ],
    "配置项清单": [
        "序号", "配置项名称", "配置项标识", "版本", "类型\n(文档/代码/其他)",
        "所属阶段WBS", "评审完成日期"
    ],
}
REPORT_ACTIVITY_PREFIXES = ("拟制质量审核报告", "拟制质量工作总结报告")
TRACKING_NAMES = ("双周进度跟踪", "阶段进度跟踪", "项目跟踪-阶段")


class WorkbookStructureError(ValueError):
    def __init__(self, diagnostics: list[Diagnostic]) -> None:
        super().__init__("Workbook structure is invalid")
        self.diagnostics = diagnostics


@dataclass(frozen=True, slots=True)
class SourceLocation:
    sheet: str
    row: int


@dataclass(frozen=True, slots=True)
class ScheduleRow:
    source: SourceLocation
    source_order: int
    wbs: str
    wbs_depth: int
    name: str
    start_date: str | None
    end_date: str | None
    linked_items: tuple[str, ...]
    linked_items_formula: str | None
    resources: tuple[str, ...]


def _json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    return value


def _date_value(value: Any) -> str | None:
    converted = _json_value(value)
    return converted if isinstance(converted, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", converted) else None


def normalize_wbs(value: Any) -> str:
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return format(Decimal(str(value)), "f").rstrip("0").rstrip(".")
    text = str(value).strip()
    return re.sub(r"\s+", "", text)


def _split_values(value: Any) -> tuple[str, ...]:
    if value is None:
        return ()
    return tuple(part.strip() for part in re.split(r"[,，;；\n]", str(value)) if part.strip())


def _headers(worksheet) -> dict[str, int]:
    return {
        str(cell.value).strip(): cell.column
        for cell in worksheet[1]
        if cell.value is not None and str(cell.value).strip()
    }


def _validate_structure(workbook) -> None:
    diagnostics: list[Diagnostic] = []
    for sheet_name, expected_headers in REQUIRED_HEADERS.items():
        if sheet_name not in workbook.sheetnames:
            diagnostics.append(
                Diagnostic(
                    "IMPORT_MISSING_SHEET", "blocker", "workbook", {"sheet": sheet_name},
                    f"缺少必需 Sheet：{sheet_name}", "请使用完整的项目相关信息表模板。",
                )
            )
            continue
        headers = _headers(workbook[sheet_name])
        for header in expected_headers:
            if header not in headers:
                diagnostics.append(
                    Diagnostic(
                        "IMPORT_MISSING_COLUMN", "blocker", "sheet",
                        {"sheet": sheet_name, "row": 1, "column_name": header},
                        f"{sheet_name} 缺少必需列：{header}", "请恢复模板中的必需列。",
                    )
                )
    if diagnostics:
        raise WorkbookStructureError(diagnostics)


def _schedule_rows(formula_ws, value_ws) -> list[ScheduleRow]:
    headers = _headers(formula_ws)
    rows: list[ScheduleRow] = []
    for row_no in range(2, formula_ws.max_row + 1):
        raw_wbs = formula_ws.cell(row_no, headers["WBS"]).value
        wbs = normalize_wbs(raw_wbs)
        name = formula_ws.cell(row_no, headers["任务名称"]).value
        if not wbs and name is None:
            continue
        resolved_linked = value_ws.cell(row_no, headers["关联配置项"]).value
        formula_linked = formula_ws.cell(row_no, headers["关联配置项"]).value
        rows.append(
            ScheduleRow(
                source=SourceLocation("项目进度计划", row_no),
                source_order=row_no - 1,
                wbs=wbs,
                wbs_depth=len(wbs.split(".")) if wbs else 0,
                name=str(name).strip() if name is not None else "",
                start_date=_date_value(value_ws.cell(row_no, headers["计划开始时间"]).value),
                end_date=_date_value(value_ws.cell(row_no, headers["计划完成时间"]).value),
                linked_items=_split_values(resolved_linked),
                linked_items_formula=(
                    str(formula_linked) if isinstance(formula_linked, str) and formula_linked.startswith("=") else None
                ),
                resources=_split_values(value_ws.cell(row_no, headers["资源名称\n(逗号分隔)"]).value),
            )
        )
    return rows


def _has_descendant(row: ScheduleRow, rows: list[ScheduleRow]) -> bool:
    prefix = f"{row.wbs}."
    return any(candidate.wbs.startswith(prefix) for candidate in rows)


def _direct_children(parent: ScheduleRow, rows: list[ScheduleRow]) -> list[ScheduleRow]:
    prefix = f"{parent.wbs}."
    return [
        row for row in rows
        if row.wbs.startswith(prefix) and row.wbs_depth == parent.wbs_depth + 1
    ]


def _project_info(formula_ws, value_ws) -> dict[str, Any]:
    headers = _headers(formula_ws)
    result: dict[str, Any] = {}
    field_map = {
        "project_name": "项目名称", "project_code": "项目代号",
        "project_leader": "项目组长", "qa_person": "质量管理员",
    }
    for key, header in field_map.items():
        result[key] = _json_value(value_ws.cell(2, headers[header]).value)
    result["source"] = {"sheet": "项目信息", "row": 2}
    return result


def _config_items(formula_ws, value_ws, diagnostics: list[Diagnostic]) -> list[dict[str, Any]]:
    headers = _headers(formula_ws)
    items: list[dict[str, Any]] = []
    identities: set[tuple[str, str]] = set()
    for row_no in range(2, formula_ws.max_row + 1):
        name = value_ws.cell(row_no, headers["配置项名称"]).value
        if name is None:
            continue
        version = value_ws.cell(row_no, headers["版本"]).value
        identity = (str(name).strip(), str(version or "").strip())
        if identity in identities:
            diagnostics.append(
                Diagnostic(
                    "IMPORT_DUPLICATE_CONFIG_ITEM", "blocker", "config_item",
                    {"sheet": "配置项清单", "row": row_no},
                    f"配置项名称和版本重复：{identity[0]} {identity[1]}", "请删除重复记录或修正版本。",
                )
            )
        identities.add(identity)
        raw_phase = formula_ws.cell(row_no, headers["所属阶段WBS"]).value
        resolved_phase = value_ws.cell(row_no, headers["所属阶段WBS"]).value
        items.append(
            {
                "source": {"sheet": "配置项清单", "row": row_no},
                "source_order": row_no - 1,
                "sequence": value_ws.cell(row_no, headers["序号"]).value,
                "name": str(name).strip(),
                "identifier": value_ws.cell(row_no, headers["配置项标识"]).value,
                "version": version,
                "item_type": value_ws.cell(row_no, headers["类型\n(文档/代码/其他)"]).value,
                "phase_wbs": normalize_wbs(resolved_phase),
                "phase_wbs_formula": raw_phase if isinstance(raw_phase, str) and raw_phase.startswith("=") else None,
                "review_date": _date_value(value_ws.cell(row_no, headers["评审完成日期"]).value),
            }
        )
    return items


def parse_project_workbook(path: Path) -> dict[str, Any]:
    formula_book = load_workbook(path, data_only=False, read_only=True, keep_links=False)
    value_book = load_workbook(path, data_only=True, read_only=True, keep_links=False)
    try:
        _validate_structure(formula_book)
        diagnostics: list[Diagnostic] = []
        rows = _schedule_rows(
            formula_book["项目进度计划"], value_book["项目进度计划"]
        )
        for row in rows:
            if not row.wbs or not re.fullmatch(r"[0-9]+(?:\.[0-9]+)*", row.wbs):
                diagnostics.append(
                    Diagnostic(
                        "IMPORT_INVALID_WBS", "blocker", "schedule", asdict(row.source),
                        f"WBS 格式无效：{row.wbs or '空'}", "请使用点分隔的数字 WBS 文本。",
                    )
                )
            if not row.name:
                diagnostics.append(
                    Diagnostic(
                        "IMPORT_MISSING_TASK_NAME", "blocker", "schedule", asdict(row.source),
                        "任务名称为空", "请补充任务名称。",
                    )
                )
            if row.linked_items_formula and not row.linked_items:
                diagnostics.append(
                    Diagnostic(
                        "IMPORT_UNRESOLVED_FORMULA", "blocker", "schedule", asdict(row.source),
                        "关联配置项公式没有可读取的缓存值", "请使用 Excel 打开并保存源文件后重新导入。",
                    )
                )
        for row in rows:
            if not _has_descendant(row, rows) and (row.start_date is None or row.end_date is None):
                diagnostics.append(
                    Diagnostic(
                        "IMPORT_MISSING_SCHEDULE_DATE", "blocker", "schedule", asdict(row.source),
                        "末级任务的计划开始或完成日期为空或格式无效", "请补充有效的 Excel 日期。",
                    )
                )

        top_rows = [row for row in rows if row.wbs_depth == 1]
        phases = [row for row in top_rows if _has_descendant(row, rows) and row.name != "公司质量验证"]
        excluded_top_level = [
            {**asdict(row), "reason": "excluded_company_validation" if row.name == "公司质量验证" else "no_children"}
            for row in top_rows if row not in phases
        ]
        quality_activities: list[dict[str, Any]] = []
        tracking_nodes: list[dict[str, Any]] = []
        for phase in phases:
            phase_children = _direct_children(phase, rows)
            for parent in phase_children:
                if parent.name == "质量保证活动":
                    for activity in _direct_children(parent, rows):
                        if len(activity.resources) > 1:
                            diagnostics.append(
                                Diagnostic(
                                    "IMPORT_MULTIPLE_QA_RESOURCES", "blocker", "activity",
                                    asdict(activity.source), "质量活动包含多名资源人员",
                                    "每个质量活动只保留一名审核人员。",
                                )
                            )
                        quality_activities.append(
                            {
                                **asdict(activity),
                                "phase_wbs": phase.wbs,
                                "parent_wbs": parent.wbs,
                                "is_report_activity": activity.name.startswith(REPORT_ACTIVITY_PREFIXES),
                            }
                        )
                if parent.name == "项目监控活动":
                    for node in _direct_children(parent, rows):
                        if any(token in node.name for token in TRACKING_NAMES):
                            tracking_nodes.append(
                                {
                                    **asdict(node),
                                    "phase_wbs": phase.wbs,
                                    "parent_wbs": parent.wbs,
                                    "node_type": "双周" if "双周" in node.name else "阶段",
                                    "node_date": node.end_date,
                                }
                            )

        config_items = _config_items(
            formula_book["配置项清单"], value_book["配置项清单"], diagnostics
        )
        for activity in quality_activities:
            classification = classify_activity(activity["name"])
            if classification.kind == "unknown":
                diagnostics.append(
                    Diagnostic(
                        "IMPORT_UNKNOWN_QUALITY_ACTIVITY", "blocker", "activity",
                        activity["source"], f"无法分类质量活动：{activity['name']}",
                        "请修正活动名称或补充分类规则。",
                    )
                )
            elif classification.kind == "product" and not activity["linked_items"]:
                diagnostics.append(
                    Diagnostic(
                        "IMPORT_PRODUCT_AUDIT_WITHOUT_CONFIG_ITEM", "blocker", "activity",
                        activity["source"], "产品审核没有关联配置项",
                        "请在关联配置项列选择工作产品。",
                    )
                )
        for item in config_items:
            if item["phase_wbs_formula"] and not item["phase_wbs"]:
                diagnostics.append(
                    Diagnostic(
                        "IMPORT_UNRESOLVED_FORMULA", "blocker", "config_item", item["source"],
                        "所属阶段WBS公式没有可读取的缓存值", "请使用 Excel 打开并保存源文件后重新导入。",
                    )
                )
        return {
            "schema_version": 1,
            "project": _project_info(formula_book["项目信息"], value_book["项目信息"]),
            "phases": [{**asdict(row), "detected_as_phase": True} for row in phases],
            "excluded_top_level": excluded_top_level,
            "tracking_nodes": tracking_nodes,
            "quality_activities": quality_activities,
            "config_items": config_items,
            "schedule_row_count": len(rows),
            "diagnostics": [diagnostic.to_dict() for diagnostic in diagnostics],
        }
    finally:
        formula_book.close()
        value_book.close()
