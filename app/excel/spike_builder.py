from __future__ import annotations

from dataclasses import dataclass
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from app.excel.external_links import add_external_workbook_link
from app.excel.formulas import (
    Formula,
    external_reference,
    internal_reference,
    write_formula,
    write_literal_cell,
)
from app.excel.templates import clear_values, copy_worksheet_template


CHECKLIST_SHEET = "软件需求规格说明检查单"
AUDIT_REPORT_SHEET = "001产品审核报告-需求规格"
PHASE_REPORT_SHEET = "阶段质量审核工作报告"
TREND_SHEET = "质量工作趋势分析"
PHASE_FILENAME = "QualityM-Step5-阶段质量审核记录.xlsx"
TREND_FILENAME = "QualityM-Step5-质量工作趋势分析.xlsx"


@dataclass(frozen=True, slots=True)
class ExcelSpikeSnapshot:
    project_name: str = "QualityM Step 5 技术原型"
    project_code: str = "STEP5"
    config_item_number: str = "0.0"
    phase_name: str = "软件需求分析"
    audit_object: str = "软件需求规格说明"
    auditor: str = "项目QA"
    responsible_person: str = "项目组长"
    audit_date: date = date(2026, 8, 14)
    issue_description: str = "需求描述未完整覆盖已确认的项目约束。"
    corrective_action: str = "补充缺失约束并完成一致性复核。"
    severity: str = "一般"


@dataclass(frozen=True, slots=True)
class ExcelSpikeArtifacts:
    phase_workbook: Path
    trend_workbook: Path


def _set_recalculation(workbook: Workbook) -> None:
    workbook.calculation.calcMode = "auto"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True


def _copy_source_sheet(source_path: Path, source_sheet: str, target, title: str) -> None:
    source_workbook = load_workbook(source_path, data_only=False, read_only=False)
    try:
        copy_worksheet_template(source_workbook[source_sheet], target)
        target.title = title
    finally:
        source_workbook.close()


def _build_checklist_sheet(workbook: Workbook, snapshot: ExcelSpikeSnapshot, template_dir: Path):
    worksheet = workbook.active
    _copy_source_sheet(
        template_dir / "工作产品检查单.xlsx",
        "软件需求规格说明检查单",
        worksheet,
        CHECKLIST_SHEET,
    )
    worksheet["A2"] = (
        f"编号： ZDPX-{snapshot.project_code}/{snapshot.config_item_number}-QA-检查单-001"
    )
    worksheet["C3"] = snapshot.project_name
    worksheet["E3"] = snapshot.phase_name
    worksheet["G3"] = snapshot.audit_date
    worksheet["K3"] = snapshot.auditor
    worksheet["C4"] = snapshot.audit_object
    worksheet["G4"] = snapshot.project_code
    worksheet["K4"] = "V1.0"

    worksheet["H7"] = "✓"
    worksheet["J7"] = snapshot.severity
    write_literal_cell(worksheet["L7"], snapshot.issue_description)
    worksheet["G8"] = "✓"
    for row in range(9, 26):
        worksheet.cell(row, 7).value = "✓"

    conclusion_validation = DataValidation(type="list", formula1='"✓,"', allow_blank=True)
    severity_validation = DataValidation(type="list", formula1='"轻微,一般,严重"')
    worksheet.add_data_validation(conclusion_validation)
    worksheet.add_data_validation(severity_validation)
    conclusion_validation.add(worksheet["H7"])
    severity_validation.add(worksheet["J7"])

    worksheet.print_area = "A1:L28"
    return worksheet


def _build_audit_report_sheet(
    workbook: Workbook,
    snapshot: ExcelSpikeSnapshot,
    template_dir: Path,
):
    worksheet = workbook.create_sheet(AUDIT_REPORT_SHEET)
    _copy_source_sheet(
        template_dir / "测控辅助业务-质量审核报告-软件需求与分析阶段.xlsx",
        "015产品审核报告-软件需规2025.8.20",
        worksheet,
        AUDIT_REPORT_SHEET,
    )
    clear_values(worksheet, "B13:K16")
    worksheet["B2"] = (
        f"编号：ZDPX-{snapshot.project_code}/{snapshot.config_item_number}-QA-审核报告-001"
    )
    worksheet["C3"] = snapshot.project_name
    worksheet["I3"] = snapshot.phase_name
    worksheet["F4"] = snapshot.audit_object
    worksheet["C5"] = snapshot.auditor
    worksheet["H5"] = snapshot.audit_date
    worksheet["C6"] = (
        f"ZDPX-{snapshot.project_code}/{snapshot.config_item_number}-QA-检查单-001"
    )
    worksheet["B8"] = "本次审核依据对应检查单执行；不符合项由检查单公式联动生成。"

    checklist_ref = lambda cell: internal_reference(CHECKLIST_SHEET, cell)
    write_formula(
        worksheet["C13"],
        Formula(f'=IF({checklist_ref("$H$7")}="✓",{checklist_ref("$L$7")},"")'),
    )
    write_formula(worksheet["B13"], Formula('=IF(C13="","",1)'))
    write_formula(
        worksheet["F13"],
        Formula(f'=IF(C13="","",{checklist_ref("$J$7")})'),
    )
    worksheet["G13"] = snapshot.responsible_person
    worksheet["H13"] = snapshot.audit_date
    worksheet["I13"] = snapshot.audit_date
    write_literal_cell(worksheet["J13"], snapshot.corrective_action)
    write_formula(worksheet["D10"], Formula('=COUNTIF(C13:C16,"?*")'))
    write_formula(worksheet["F10"], Formula('=COUNTIF(F13:F16,"轻微")'))
    write_formula(worksheet["H10"], Formula('=COUNTIF(F13:F16,"一般")'))
    write_formula(worksheet["J10"], Formula('=COUNTIF(F13:F16,"严重")'))
    worksheet["B18"] = "纠正措施已实施，验证结果以导出快照为准。\n\n验证人/日期："
    worksheet.print_area = "B1:K18"
    return worksheet


def _build_phase_report_sheet(
    workbook: Workbook,
    snapshot: ExcelSpikeSnapshot,
    template_dir: Path,
):
    worksheet = workbook.create_sheet(PHASE_REPORT_SHEET)
    _copy_source_sheet(
        template_dir / "测控辅助业务-质量审核报告-软件需求与分析阶段.xlsx",
        "质量审核工作报告2025.8.25",
        worksheet,
        PHASE_REPORT_SHEET,
    )
    for row in worksheet.iter_rows():
        for cell in row:
            if getattr(cell, "data_type", None) == "f":
                cell.value = None
    clear_values(worksheet, "A6:N86")

    worksheet["A2"] = (
        f"编号：ZDPX-{snapshot.project_code}/{snapshot.config_item_number}-QA-质量报告-001"
    )
    worksheet["B3"] = snapshot.project_name
    worksheet["B4"] = snapshot.phase_name
    worksheet["F4"] = snapshot.audit_date
    worksheet["K4"] = snapshot.auditor
    worksheet["A6"] = "审核过程/产品名称"
    write_formula(
        worksheet["C6"],
        Formula(f"={internal_reference(AUDIT_REPORT_SHEET, '$F$4')}"),
    )
    worksheet["K6"] = "审核时间"
    write_formula(
        worksheet["L6"],
        Formula(f"={internal_reference(AUDIT_REPORT_SHEET, '$H$5')}"),
    )
    worksheet["A7"] = "不符合项列表"
    worksheet["J7"] = "NC类型"
    worksheet["K7"] = "严重程度"
    worksheet["L7"] = "是否关闭"
    write_formula(
        worksheet["A8"],
        Formula(f'=IF({internal_reference(AUDIT_REPORT_SHEET, "$B$13")}="","",{internal_reference(AUDIT_REPORT_SHEET, "$B$13")})'),
    )
    write_formula(
        worksheet["B8"],
        Formula(f'=IF({internal_reference(AUDIT_REPORT_SHEET, "$C$13")}="","",{internal_reference(AUDIT_REPORT_SHEET, "$C$13")})'),
    )
    worksheet["J8"] = "工作产品"
    write_formula(
        worksheet["K8"],
        Formula(f'=IF($B$8="","",{internal_reference(AUDIT_REPORT_SHEET, "$F$13")})'),
    )
    worksheet["L8"] = "是"
    for hidden_row in range(9, 87):
        worksheet.row_dimensions[hidden_row].hidden = True

    worksheet["C87"] = "本阶段技术原型包含一个产品审核实例。"
    worksheet["C88"] = snapshot.audit_date
    worksheet["I88"] = snapshot.audit_date
    write_formula(
        worksheet["A91"],
        Formula(f"={internal_reference(AUDIT_REPORT_SHEET, '$D$10')}"),
    )
    worksheet["C92"] = 0
    write_formula(worksheet["E92"], Formula('=IFERROR((A91-C92)/A91,0)'))
    write_formula(worksheet["I92"], Formula('=COUNTIF($K$8,"轻微")'))
    write_formula(worksheet["K92"], Formula('=COUNTIF($K$8,"一般")'))
    write_formula(worksheet["L92"], Formula('=COUNTIF($K$8,"严重")'))
    write_formula(worksheet["I93"], Formula('=IFERROR(I92/A91,0)'))
    write_formula(worksheet["K93"], Formula('=IFERROR(K92/A91,0)'))
    write_formula(worksheet["L93"], Formula('=IFERROR(L92/A91,0)'))
    worksheet["A96"] = 0
    worksheet["C97"] = 0
    worksheet["E97"] = 0
    worksheet["I97"] = 0
    worksheet["K97"] = 0
    worksheet["L97"] = 0
    for column in "CDEFGHIJ":
        worksheet[f"{column}101"] = 0
    write_formula(worksheet["K101"], Formula("=$A$91"))
    write_formula(worksheet["L101"], Formula("=SUM(C101:K101)"))
    for column in "CDEFGHIJK":
        write_formula(
            worksheet[f"{column}102"],
            Formula(f"=IFERROR({column}101/$L$101,0)"),
        )
    worksheet["A104"] = "本阶段原型统计由检查单、单次审核报告和阶段报告公式逐级联动。"
    worksheet.print_area = "A1:L104"
    return worksheet


def _build_trend_workbook(snapshot: ExcelSpikeSnapshot, phase_filename: str) -> Workbook:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = TREND_SHEET
    worksheet.sheet_view.showGridLines = False
    _set_recalculation(workbook)
    link_index = add_external_workbook_link(
        workbook,
        target=phase_filename,
        sheet_names=(PHASE_REPORT_SHEET,),
    )

    worksheet.merge_cells("A1:E1")
    worksheet["A1"] = "项目质量工作趋势分析（Step 5 技术原型）"
    worksheet["A1"].font = Font(name="微软雅黑", size=16, bold=True, color="FFFFFF")
    worksheet["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    worksheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    worksheet.row_dimensions[1].height = 30
    headers = ("阶段", "不符合项总数", "轻微", "一般", "严重")
    for column, value in enumerate(headers, 1):
        cell = worksheet.cell(3, column, value)
        cell.font = Font(name="微软雅黑", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="5B9BD5")
        cell.alignment = Alignment(horizontal="center")
    worksheet["A4"] = snapshot.phase_name
    external = lambda cell: external_reference(link_index, PHASE_REPORT_SHEET, cell)
    write_formula(worksheet["B4"], Formula(f"={external('$L$101')}"))
    write_formula(worksheet["C4"], Formula(f"={external('$I$92')}+{external('$I$97')}"))
    write_formula(worksheet["D4"], Formula(f"={external('$K$92')}+{external('$K$97')}"))
    write_formula(worksheet["E4"], Formula(f"={external('$L$92')}+{external('$L$97')}"))
    worksheet["A5"] = "合计"
    for column in range(2, 6):
        letter = get_column_letter(column)
        write_formula(worksheet.cell(5, column), Formula(f"=SUM({letter}4:{letter}4)"))
    worksheet["A7"] = "统计结果显示"
    worksheet["A8"] = "阶段不符合项数量随阶段工作簿更新链接后自动刷新。"
    worksheet.merge_cells("A8:E8")
    worksheet["A8"].alignment = Alignment(wrap_text=True)
    worksheet.column_dimensions["A"].width = 22
    for column in "BCDE":
        worksheet.column_dimensions[column].width = 16
    worksheet.freeze_panes = "A4"
    worksheet.print_area = "A1:J20"

    chart = LineChart()
    chart.title = "阶段不符合项数量趋势图"
    chart.y_axis.title = "数量"
    chart.x_axis.title = "阶段"
    chart.add_data(Reference(worksheet, min_col=2, max_col=5, min_row=3, max_row=4), titles_from_data=True)
    chart.set_categories(Reference(worksheet, min_col=1, min_row=4, max_row=4))
    chart.height = 8
    chart.width = 14
    worksheet.add_chart(chart, "G2")
    return workbook


def build_excel_linkage_spike(
    snapshot: ExcelSpikeSnapshot,
    *,
    output_dir: Path,
    template_dir: Path,
) -> ExcelSpikeArtifacts:
    output_dir = output_dir.resolve()
    template_dir = template_dir.resolve()
    required_templates = (
        template_dir / "工作产品检查单.xlsx",
        template_dir / "测控辅助业务-质量审核报告-软件需求与分析阶段.xlsx",
    )
    if any(not path.is_file() for path in required_templates):
        raise FileNotFoundError("Step 5 authoritative Excel templates are incomplete")
    output_dir.mkdir(parents=True, exist_ok=True)

    phase_workbook = Workbook()
    _set_recalculation(phase_workbook)
    _build_checklist_sheet(phase_workbook, snapshot, template_dir)
    _build_audit_report_sheet(phase_workbook, snapshot, template_dir)
    _build_phase_report_sheet(phase_workbook, snapshot, template_dir)
    phase_path = output_dir / PHASE_FILENAME
    phase_workbook.save(phase_path)
    phase_workbook.close()

    trend_workbook = _build_trend_workbook(snapshot, phase_path.name)
    trend_path = output_dir / TREND_FILENAME
    trend_workbook.save(trend_path)
    trend_workbook.close()
    return ExcelSpikeArtifacts(phase_workbook=phase_path, trend_workbook=trend_path)
