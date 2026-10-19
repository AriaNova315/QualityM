from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Alignment, Font, PatternFill

from app.excel.formulas import write_literal_cell


SHEET_NAME = "模板库"
HEADERS = (
    "检查单稳定键", "检查单名称", "检查单版本", "审核类型", "检查单启用", "检查单顺序",
    "分组稳定键", "分组名称", "分组版本", "分组启用", "分组顺序",
    "检查项稳定键", "检查项显示号", "检查项内容", "检查项星标", "检查项版本", "检查项启用", "检查项顺序",
    "问题稳定键", "问题版本", "问题标题", "问题描述", "整改建议", "默认严重度", "占位符", "默认问题", "问题启用", "问题顺序",
)


def build_template_library(rows: tuple[dict, ...], path: Path) -> Path:
    book = Workbook(); ws = book.active; ws.title = SHEET_NAME
    ws.append(HEADERS)
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.fill = fill; cell.font = Font(name="宋体", color="FFFFFF", bold=True); cell.alignment = Alignment(wrap_text=True)
    ws.row_dimensions[1].height = 30
    for row_index, row in enumerate(rows, start=2):
        values = [row[key] for key in (
            "checklist_key", "checklist_name", "checklist_version", "audit_kind", "checklist_enabled", "checklist_order",
            "group_key", "group_name", "group_version", "group_enabled", "group_order",
            "item_key", "display_no", "content", "starred", "item_version", "item_enabled", "item_order",
            "issue_key", "issue_version", "title", "description", "corrective_suggestion", "default_severity", "placeholders", "is_default", "issue_enabled", "issue_order",
        )]
        for column_index, value in enumerate(values, start=1):
            cell = ws.cell(row_index, column_index)
            if isinstance(value, str):
                write_literal_cell(cell, value)
            else:
                cell.value = value
            cell.font = Font(name="宋体", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=column_index in {2, 7, 13, 20, 21, 22})
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    widths = (26, 24, 14, 12, 12, 12, 26, 24, 14, 12, 12, 28, 14, 48, 12, 14, 12, 12, 44, 14, 28, 54, 54, 14, 28, 12, 12, 12)
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    for column, values in {
        "D": '"product,process"', "E": '"0,1"', "J": '"0,1"', "O": '"0,1"',
        "Q": '"0,1"', "X": '"轻微,一般,严重"', "Z": '"0,1"', "AA": '"0,1"',
    }.items():
        validation = DataValidation(type="list", formula1=values, allow_blank=False)
        validation.error = "请选择下拉列表中的有效值"
        validation.errorTitle = "模板数据无效"
        ws.add_data_validation(validation)
        validation.add(f"{column}2:{column}1048576")
    ws.sheet_view.showGridLines = False; ws.sheet_view.zoomScale = 80
    path.parent.mkdir(parents=True, exist_ok=True); book.save(path); book.close(); return path


def read_template_library(path: Path) -> tuple[dict, ...]:
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        if book.sheetnames != [SHEET_NAME]: raise ValueError("模板库 Excel 必须只包含“模板库”工作表")
        ws = book[SHEET_NAME]
        headers = tuple(cell.value for cell in ws[1])
        if headers != HEADERS: raise ValueError("模板库 Excel 列结构不匹配")
        rows = []
        keys = (
            "checklist_key", "checklist_name", "checklist_version", "audit_kind", "checklist_enabled", "checklist_order",
            "group_key", "group_name", "group_version", "group_enabled", "group_order",
            "item_key", "display_no", "content", "starred", "item_version", "item_enabled", "item_order",
            "issue_key", "issue_version", "title", "description", "corrective_suggestion", "default_severity", "placeholders", "is_default", "issue_enabled", "issue_order",
        )
        for values in ws.iter_rows(min_row=2, values_only=True):
            if all(value is None for value in values): continue
            if len(values) != len(keys): raise ValueError("模板库 Excel 存在不完整行")
            rows.append(dict(zip(keys, values, strict=True)))
        if not rows: raise ValueError("模板库 Excel 不包含模板数据")
        return tuple(rows)
    finally: book.close()
