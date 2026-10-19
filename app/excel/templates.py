from __future__ import annotations

from copy import copy

from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter


def copy_cell_format(source_cell, target_cell) -> None:
    """Copy one cell's presentation without sharing mutable style objects."""

    if source_cell.has_style:
        target_cell.font = copy(source_cell.font)
        target_cell.fill = copy(source_cell.fill)
        target_cell.border = copy(source_cell.border)
        target_cell.alignment = copy(source_cell.alignment)
        target_cell.protection = copy(source_cell.protection)
        target_cell.number_format = source_cell.number_format


def copy_worksheet_template(source, target) -> None:
    """Copy a worksheet's visible template without sharing mutable styles."""

    for row in source.iter_rows():
        for source_cell in row:
            if isinstance(source_cell, MergedCell):
                continue
            target_cell = target[source_cell.coordinate]
            target_cell.value = source_cell.value
            copy_cell_format(source_cell, target_cell)
            if source_cell.hyperlink:
                target_cell._hyperlink = copy(source_cell.hyperlink)
            if source_cell.comment:
                target_cell.comment = copy(source_cell.comment)

    for key, source_dimension in source.row_dimensions.items():
        target_dimension = target.row_dimensions[key]
        target_dimension.height = source_dimension.height
        target_dimension.hidden = source_dimension.hidden
        target_dimension.outlineLevel = source_dimension.outlineLevel
        target_dimension.collapsed = source_dimension.collapsed
    for key, source_dimension in source.column_dimensions.items():
        target_dimension = target.column_dimensions[key]
        target_dimension.width = source_dimension.width
        target_dimension.hidden = source_dimension.hidden
        target_dimension.bestFit = source_dimension.bestFit
        target_dimension.outlineLevel = source_dimension.outlineLevel
        target_dimension.collapsed = source_dimension.collapsed
        target_dimension.min = source_dimension.min
        target_dimension.max = source_dimension.max
    for merged_range in source.merged_cells.ranges:
        target.merge_cells(str(merged_range))

    target.sheet_format = copy(source.sheet_format)
    target.sheet_properties = copy(source.sheet_properties)
    target.page_margins = copy(source.page_margins)
    target.page_setup = copy(source.page_setup)
    target.print_options = copy(source.print_options)
    target.freeze_panes = source.freeze_panes
    target.sheet_view.showGridLines = source.sheet_view.showGridLines


def copy_worksheet_columns(source, target) -> None:
    """Copy column dimensions and page-level presentation from a source sheet."""

    for index in range(1, source.max_column + 1):
        key = get_column_letter(index)
        source_dimension = source.column_dimensions[key]
        target_dimension = target.column_dimensions[key]
        target_dimension.width = source_dimension.width
        target_dimension.hidden = source_dimension.hidden
        target_dimension.bestFit = source_dimension.bestFit
        target_dimension.outlineLevel = source_dimension.outlineLevel
        target_dimension.collapsed = source_dimension.collapsed
    target.sheet_format = copy(source.sheet_format)
    target.sheet_properties = copy(source.sheet_properties)
    target.page_margins = copy(source.page_margins)
    target.page_setup = copy(source.page_setup)
    target.print_options = copy(source.print_options)
    target.sheet_view.showGridLines = source.sheet_view.showGridLines


def copy_worksheet_rows(source, target, source_start: int, source_end: int, target_start: int) -> None:
    """Copy a contiguous template block, translating cells and merged regions."""

    offset = target_start - source_start
    for source_row in range(source_start, source_end + 1):
        target_row = source_row + offset
        source_dimension = source.row_dimensions[source_row]
        target_dimension = target.row_dimensions[target_row]
        target_dimension.height = source_dimension.height
        target_dimension.hidden = source_dimension.hidden
        target_dimension.outlineLevel = source_dimension.outlineLevel
        target_dimension.collapsed = source_dimension.collapsed
        for source_cell in source[source_row]:
            if isinstance(source_cell, MergedCell):
                continue
            target_cell = target.cell(target_row, source_cell.column)
            target_cell.value = source_cell.value
            copy_cell_format(source_cell, target_cell)

    for merged_range in source.merged_cells.ranges:
        if merged_range.min_row < source_start or merged_range.max_row > source_end:
            continue
        target.merge_cells(
            start_row=merged_range.min_row + offset,
            start_column=merged_range.min_col,
            end_row=merged_range.max_row + offset,
            end_column=merged_range.max_col,
        )


def clear_values(worksheet, cell_range: str) -> None:
    for row in worksheet[cell_range]:
        for cell in row:
            if not isinstance(cell, MergedCell):
                cell.value = None
