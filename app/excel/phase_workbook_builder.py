from __future__ import annotations

from collections import Counter
from copy import copy
from dataclasses import dataclass
from datetime import date, timedelta
from math import ceil
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.application.export_snapshot_service import (
    ExportAuditSnapshot,
    ExportIssueSnapshot,
    ExportPhaseSnapshot,
    ExportSnapshot,
)
from app.excel.formula_cache import write_formula_caches
from app.excel.formulas import Formula, internal_reference, write_formula, write_literal_cell
from app.excel.sheet_names import safe_filename_component, safe_sheet_name
from app.excel.templates import (
    copy_cell_format,
    copy_worksheet_columns,
    copy_worksheet_rows,
    copy_worksheet_template,
)


DIRECTORY_SHEET = "审核目录"
PHASE_REPORT_SHEET = "阶段质量审核工作报告"
CALCULATION_SHEET = "_计算"

_REFERENCE_DIR = Path(__file__).resolve().parents[2] / "temp"
_REPORT_REFERENCE = _REFERENCE_DIR / "测控辅助业务-质量审核报告-软件需求与分析阶段.xlsx"
_PRODUCT_CHECKLIST_REFERENCE = _REFERENCE_DIR / "工作产品检查单.xlsx"
_PROCESS_CHECKLIST_REFERENCE = _REFERENCE_DIR / "过程检查单.xlsx"
_AUDIT_REPORT_TEMPLATE_SHEET = "005产品审核报告-测量与分析计划2025.7.15"
_PHASE_REPORT_TEMPLATE_SHEET = "质量审核工作报告2025.8.25"
_PHASE_VISIBLE_CATEGORIES = ("RDM", "PP", "PMC", "CM", "MT", "MPM", "VV", "II", "工作产品")
_TREND_CATEGORIES = ("RDM", "PP", "PMC", "CM", "MT", "MPM", "VV", "LD", "II", "工作产品")

_THIN = Side(style="thin", color="000000")
_THICK = Side(style="thick", color="000000")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_TITLE_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
_GROUP_FILL = PatternFill("solid", fgColor="E2F0D9")
_SUMMARY_FILL = PatternFill("solid", fgColor="FFF2CC")


@dataclass(frozen=True, slots=True)
class PhaseWorkbookArtifact:
    phase_id: int
    phase_name: str
    path: Path
    audit_count: int
    issue_count: int


@dataclass(frozen=True, slots=True)
class _AuditSheetInfo:
    audit: ExportAuditSnapshot
    checklist_sheet: str
    report_sheet: str
    report_issue_rows: tuple[int, ...]


def _set_recalculation(workbook: Workbook) -> None:
    workbook.calculation.calcMode = "auto"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True


def _style_range(worksheet, cell_range: str, *, fill=None, bold: bool = False) -> None:
    for row in worksheet[cell_range]:
        for cell in row:
            cell.border = _BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if fill is not None:
                cell.fill = fill
            if bold:
                cell.font = Font(name="宋体", size=10, bold=True)


def _merge_write(worksheet, cell_range: str, value: object, **style) -> None:
    worksheet.merge_cells(cell_range)
    cell = worksheet[cell_range.split(":", 1)[0]]
    write_literal_cell(cell, value)
    _style_range(worksheet, cell_range, **style)


def _write_formula(cell, expression: str) -> None:
    write_formula(cell, Formula(expression))


def _wrapped_height(*values: object, chars_per_line: int, max_lines: int = 4) -> float:
    longest = max((len(str(value or "")) for value in values), default=0)
    lines = max(1, min(max_lines, ceil(longest / chars_per_line)))
    return max(18.0, lines * 15.0)


def _template_height(
    *values: object,
    chars_per_line: int,
    minimum: float,
    line_height: float = 19.0,
) -> float:
    """Estimate merged-cell height while preserving the reference minimum."""

    longest = max((len(str(value or "")) for value in values), default=0)
    lines = max(1, ceil(longest / chars_per_line))
    return max(minimum, lines * line_height)


def _clear_row_values(worksheet, start_row: int, end_row: int, max_column: int = 12) -> None:
    for row in worksheet.iter_rows(
        min_row=start_row,
        max_row=end_row,
        min_col=1,
        max_col=max_column,
    ):
        for cell in row:
            if not isinstance(cell, MergedCell):
                cell.value = None


def _copy_row_format(source, source_row: int, target, target_row: int, max_column: int) -> None:
    target.row_dimensions[target_row].height = source.row_dimensions[source_row].height
    for column in range(1, max_column + 1):
        copy_cell_format(source.cell(source_row, column), target.cell(target_row, column))


def _apply_outer_border(
    worksheet,
    *,
    min_row: int,
    min_column: int,
    max_row: int,
    max_column: int,
    side: Side,
) -> None:
    """Apply one continuous outer border without changing interior borders."""

    for row in range(min_row, max_row + 1):
        for column in range(min_column, max_column + 1):
            if row not in {min_row, max_row} and column not in {min_column, max_column}:
                continue
            cell = worksheet.cell(row, column)
            border = copy(cell.border)
            if row == min_row:
                border.top = side
            if row == max_row:
                border.bottom = side
            if column == min_column:
                border.left = side
            if column == max_column:
                border.right = side
            cell.border = border

    # Merged-cell edge styles are reconstructed from the top-left anchor when
    # openpyxl saves and reopens a workbook. Mirror perimeter sides onto that
    # anchor so the continuous outline survives round-tripping through Excel.
    for merged_range in worksheet.merged_cells.ranges:
        touches_top = merged_range.min_row == min_row
        touches_bottom = merged_range.max_row == max_row
        touches_left = merged_range.min_col == min_column
        touches_right = merged_range.max_col == max_column
        if not any((touches_top, touches_bottom, touches_left, touches_right)):
            continue
        anchor = worksheet.cell(merged_range.min_row, merged_range.min_col)
        border = copy(anchor.border)
        if touches_top:
            border.top = side
        if touches_bottom:
            border.bottom = side
        if touches_left:
            border.left = side
        if touches_right:
            border.right = side
        anchor.border = border


def _merged_anchor_cell(worksheet, row: int, column: int):
    cell = worksheet.cell(row, column)
    if not isinstance(cell, MergedCell):
        return cell
    for merged_range in worksheet.merged_cells.ranges:
        if (
            merged_range.min_row <= row <= merged_range.max_row
            and merged_range.min_col <= column <= merged_range.max_col
        ):
            return worksheet.cell(merged_range.min_row, merged_range.min_col)
    raise ValueError(f"Merged cell {cell.coordinate} has no containing range")


def _iso_date(value: str) -> date:
    return date.fromisoformat(value)


def _audit_scope_text(audits: tuple[ExportAuditSnapshot, ...], period_label: str) -> str:
    processes = list(
        dict.fromkeys(
            _audit_object(audit)
            for audit in audits
            if audit.audit_kind != "product"
        )
    )
    products = list(
        dict.fromkeys(
            _audit_object(audit)
            for audit in audits
            if audit.audit_kind == "product"
        )
    )
    lines = []
    if processes:
        lines.append(f"本{period_label}审核的实践域包括：{'、'.join(processes)}。")
    if products:
        lines.append(f"审核的工作产品包括：{'、'.join(products)}。")
    return "\n".join(lines) or f"本{period_label}无审核活动。"


def _analysis_text(audits: tuple[ExportAuditSnapshot, ...], period_label: str) -> str:
    issues = tuple(issue for audit in audits for issue in _all_issues(audit))
    if not issues:
        return f"本{period_label}未发现不符合项。"
    categories = Counter(
        "工作产品" if audit.audit_kind == "product" else (audit.process_code or "过程")
        for audit in audits
        for _issue in _all_issues(audit)
    )
    maximum = max(categories.values())
    dominant = "、".join(key for key, count in categories.items() if count == maximum)
    severe = sum(issue.severity == "严重" for issue in issues)
    unclosed = sum(not issue.confirmed for issue in issues)
    severity_text = f"其中严重不符合项{severe}个" if severe else "无严重不符合项"
    return (
        f"本{period_label}共发现{len(issues)}个不符合项，主要分布在{dominant}，{severity_text}。"
        f"截止报告期，已关闭{len(issues) - unclosed}项，未关闭{unclosed}项。"
    )


def _audit_overview(audit: ExportAuditSnapshot) -> str:
    issues = _all_issues(audit)
    if audit.audit_kind == "product":
        body = (
            f"本次审核的产品为{_audit_object(audit)}，审核内容包括：\n"
            "1）编写的规范性；\n"
            "2）与其它相关文档的一致性；\n"
            "3）内容的完整性和正确性；\n"
            "4）上下文的一致性；\n"
            "5）对模板的符合性。"
        )
    else:
        names = list(dict.fromkeys(group.name for group in audit.groups if group.name))
        points = "\n".join(f"{index}）{name}；" for index, name in enumerate(names, start=1))
        body = "本次审核内容包括：" + (f"\n{points}" if points else "")
    conclusion = (
        f"经审核，发现{len(issues)}个不符合项。"
        if issues
        else "经审核，未发现不符合项。"
    )
    return f"审核内容概述：\n{body}\n{conclusion}"


def _audit_object(audit: ExportAuditSnapshot) -> str:
    return audit.config_item_name or audit.activity_name


def _audit_state_label(value: str) -> str:
    return {"draft": "未确认", "confirmed": "已确认"}.get(value, value)


def _all_issues(audit: ExportAuditSnapshot) -> tuple[ExportIssueSnapshot, ...]:
    return tuple(
        issue
        for group in audit.groups
        for item in group.items
        for issue in item.issues
    )


def _build_directory(
    workbook: Workbook,
    snapshot: ExportSnapshot,
    phase: ExportPhaseSnapshot,
    infos: tuple[_AuditSheetInfo, ...],
) -> None:
    worksheet = workbook.active
    worksheet.title = DIRECTORY_SHEET
    worksheet.sheet_view.showGridLines = False
    worksheet.column_dimensions["A"].width = 8
    worksheet.column_dimensions["B"].width = 34
    worksheet.column_dimensions["C"].width = 18
    worksheet.column_dimensions["D"].width = 18
    _merge_write(worksheet, "A1:D1", "阶段质量审核记录目录", fill=_TITLE_FILL, bold=True)
    worksheet["A1"].font = Font(name="宋体", size=18, bold=True, color="FFFFFF")
    _merge_write(worksheet, "A2:B2", f"项目：{snapshot.project_name}")
    _merge_write(worksheet, "C2:D2", f"阶段：{phase.name}")
    headers = ("流水号", "审核对象", "检查单", "审核报告")
    for column, value in enumerate(headers, start=1):
        worksheet.cell(4, column, value)
    _style_range(worksheet, "A4:D4", fill=_HEADER_FILL, bold=True)
    for row, info in enumerate(infos, start=5):
        worksheet.cell(row, 1, info.audit.sequence_no)
        write_literal_cell(worksheet.cell(row, 2), _audit_object(info.audit))
        checklist = worksheet.cell(row, 3, "打开检查单")
        checklist.hyperlink = f"#'{info.checklist_sheet}'!A1"
        checklist.style = "Hyperlink"
        report = worksheet.cell(row, 4, "打开审核报告")
        report.hyperlink = f"#'{info.report_sheet}'!A1"
        report.style = "Hyperlink"
        _style_range(worksheet, f"A{row}:D{row}")
        worksheet.row_dimensions[row].height = _wrapped_height(
            _audit_object(info.audit), chars_per_line=20, max_lines=2
        )
    worksheet.freeze_panes = "A5"
    worksheet.auto_filter.ref = f"A4:D{max(4, 4 + len(infos))}"
    worksheet.print_area = f"A1:D{max(5, 4 + len(infos))}"
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_setup.paperSize = worksheet.PAPERSIZE_A4
    worksheet.page_setup.orientation = worksheet.ORIENTATION_PORTRAIT
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.print_options.horizontalCentered = True


def _build_checklist_generic(
    worksheet,
    snapshot: ExportSnapshot,
    phase: ExportPhaseSnapshot,
    audit: ExportAuditSnapshot,
) -> dict[int, int]:
    worksheet.sheet_view.showGridLines = False
    widths = (8, 18, 18, 18, 12, 12, 8, 8, 10, 12, 12, 28)
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width
    _merge_write(worksheet, "A1:L1", audit.checklist_name, fill=_TITLE_FILL, bold=True)
    worksheet["A1"].font = Font(name="宋体", size=16, bold=True, color="FFFFFF")
    _merge_write(
        worksheet,
        "A2:L2",
        f"编号：ZDPX-{snapshot.project_code}/{snapshot.config_item_number}-QA-检查单-{audit.sequence_no:03d}",
    )
    _merge_write(worksheet, "A3:B3", "项目名称", fill=_HEADER_FILL, bold=True)
    _merge_write(worksheet, "C3:E3", snapshot.project_name)
    _merge_write(worksheet, "F3:G3", "当前阶段", fill=_HEADER_FILL, bold=True)
    _merge_write(worksheet, "H3:I3", phase.name)
    _merge_write(worksheet, "J3:K3", "项目QA", fill=_HEADER_FILL, bold=True)
    write_literal_cell(worksheet["L3"], audit.auditor or snapshot.qa_person)
    _style_range(worksheet, "L3:L3")
    _merge_write(worksheet, "A4:B4", "审核对象", fill=_HEADER_FILL, bold=True)
    _merge_write(worksheet, "C4:E4", _audit_object(audit))
    _merge_write(worksheet, "F4:G4", "产品标识", fill=_HEADER_FILL, bold=True)
    _merge_write(worksheet, "H4:I4", audit.config_item_identifier or "—")
    _merge_write(worksheet, "J4:K4", "产品版本", fill=_HEADER_FILL, bold=True)
    write_literal_cell(worksheet["L4"], audit.config_item_version or "—")
    _style_range(worksheet, "L4:L4")
    headers = ("序号", "审核内容", "", "", "", "", "是", "否", "不适用", "严重程度", "", "备注")
    for column, value in enumerate(headers, start=1):
        worksheet.cell(6, column, value)
    _style_range(worksheet, "A6:L6", fill=_HEADER_FILL, bold=True)

    row = 7
    for group in audit.groups:
        _merge_write(worksheet, f"A{row}:L{row}", group.name, fill=_GROUP_FILL, bold=True)
        row += 1
        for item in group.items:
            display_no = f"{item.display_no}{'*' if item.starred and not str(item.display_no).endswith('*') else ''}"
            write_literal_cell(worksheet.cell(row, 1), display_no)
            _merge_write(worksheet, f"B{row}:F{row}", item.content)
            conclusion_column = {"符合": 7, "不符合": 8, "不适用": 9}.get(item.conclusion)
            if conclusion_column is not None:
                worksheet.cell(row, conclusion_column, "✓")
            severities = "、".join(dict.fromkeys(issue.severity for issue in item.issues))
            write_literal_cell(worksheet.cell(row, 10), severities)
            issue_note = "；".join(issue.title for issue in item.issues)
            _merge_write(worksheet, f"K{row}:L{row}", issue_note)
            _style_range(worksheet, f"A{row}:J{row}")
            worksheet.row_dimensions[row].height = _wrapped_height(
                item.content, issue_note, chars_per_line=42, max_lines=3
            )
            row += 1

    row += 1
    _merge_write(worksheet, f"A{row}:L{row}", "不符合项问题明细", fill=_SUMMARY_FILL, bold=True)
    row += 1
    detail_headers = (
        "问题ID", "检查项", "问题描述", "", "", "严重程度",
        "责任人", "提交日期", "计划关闭日期", "纠正措施", "", "关闭",
    )
    for column, value in enumerate(detail_headers, start=1):
        worksheet.cell(row, column, value)
    _style_range(worksheet, f"A{row}:L{row}", fill=_HEADER_FILL, bold=True)
    issue_rows: dict[int, int] = {}
    row += 1
    for group in audit.groups:
        for item in group.items:
            for issue in item.issues:
                issue_rows[issue.id] = row
                worksheet.cell(row, 1, issue.id)
                display_no = f"{item.display_no}{'*' if item.starred and not str(item.display_no).endswith('*') else ''}"
                write_literal_cell(worksheet.cell(row, 2), display_no)
                _merge_write(worksheet, f"C{row}:E{row}", issue.description)
                write_literal_cell(worksheet.cell(row, 6), issue.severity)
                write_literal_cell(worksheet.cell(row, 7), issue.responsible_person)
                write_literal_cell(worksheet.cell(row, 8), issue.submitted_date)
                write_literal_cell(worksheet.cell(row, 9), issue.planned_close_date)
                corrective_action = issue.corrective_action
                if issue.suggested_closed and not issue.confirmed and corrective_action:
                    corrective_action = f"待确认的建议值：{corrective_action}"
                _merge_write(worksheet, f"J{row}:K{row}", corrective_action)
                worksheet.cell(row, 12, "是" if issue.confirmed else "否")
                _style_range(worksheet, f"A{row}:B{row}")
                _style_range(worksheet, f"F{row}:I{row}")
                _style_range(worksheet, f"L{row}:L{row}")
                worksheet.row_dimensions[row].height = max(
                    _wrapped_height(issue.description, chars_per_line=28, max_lines=3),
                    _wrapped_height(corrective_action, chars_per_line=18, max_lines=3),
                )
                row += 1
    if not issue_rows:
        _merge_write(worksheet, f"A{row}:L{row}", "本次审核未发现不符合项。")
        row += 1
    worksheet.freeze_panes = "A7"
    worksheet.print_title_rows = "1:6"
    worksheet.print_area = f"A1:L{row - 1}"
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_setup.paperSize = worksheet.PAPERSIZE_A4
    worksheet.page_setup.orientation = worksheet.ORIENTATION_LANDSCAPE
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.print_options.horizontalCentered = True
    return issue_rows


def _build_checklist(
    worksheet,
    snapshot: ExportSnapshot,
    phase: ExportPhaseSnapshot,
    audit: ExportAuditSnapshot,
    template_worksheet=None,
) -> dict[int, int]:
    """Populate an authoritative checklist layout without appending extra sections."""

    if template_worksheet is None:
        return _build_checklist_generic(worksheet, snapshot, phase, audit)

    original_title = worksheet.title
    copy_worksheet_template(template_worksheet, worksheet)
    worksheet.title = original_title
    worksheet.sheet_view.showGridLines = template_worksheet.sheet_view.showGridLines
    worksheet["A2"] = (
        f"编号：ZDPX-{snapshot.project_code}/{snapshot.config_item_number}"
        f"-QA-检查单-{audit.sequence_no:03d}"
    )

    if audit.audit_kind == "product":
        write_literal_cell(worksheet["C3"], snapshot.project_name)
        write_literal_cell(worksheet["E3"], phase.name)
        write_literal_cell(worksheet["G3"], audit.audit_date)
        write_literal_cell(worksheet["K3"], audit.auditor or snapshot.qa_person)
        write_literal_cell(worksheet["C4"], _audit_object(audit))
        write_literal_cell(worksheet["G4"], audit.config_item_identifier or "—")
        write_literal_cell(worksheet["K4"], audit.config_item_version or "—")
        conclusion_columns = {"符合": 7, "不符合": 8, "不适用": 9}
        severity_column, note_column = 10, 12
    else:
        write_literal_cell(worksheet["C3"], snapshot.project_name)
        write_literal_cell(worksheet["E3"], phase.name)
        write_literal_cell(worksheet["J3"], audit.audit_date)
        write_literal_cell(worksheet["M3"], audit.auditor or snapshot.qa_person)
        conclusion_columns = {"符合": 6, "不符合": 8, "不适用": 9}
        severity_column, note_column = 10, 12

    issue_rows: dict[int, int] = {}
    for group in audit.groups:
        if group.source_row is not None:
            write_literal_cell(worksheet.cell(group.source_row, 1), group.name)
        for item in group.items:
            if item.source_row is None:
                continue
            row = item.source_row
            display_no = f"{item.display_no}{'*' if item.starred and not str(item.display_no).endswith('*') else ''}"
            write_literal_cell(worksheet.cell(row, 1), display_no)
            write_literal_cell(worksheet.cell(row, 2), item.content)
            for column in set(conclusion_columns.values()):
                _merged_anchor_cell(worksheet, row, column).value = None
            _merged_anchor_cell(worksheet, row, severity_column).value = None
            _merged_anchor_cell(worksheet, row, note_column).value = None
            conclusion_column = conclusion_columns.get(item.conclusion)
            if conclusion_column is not None:
                _merged_anchor_cell(worksheet, row, conclusion_column).value = "✓"
            severities = "、".join(dict.fromkeys(issue.severity for issue in item.issues))
            descriptions = "；".join(issue.description for issue in item.issues)
            write_literal_cell(_merged_anchor_cell(worksheet, row, severity_column), severities)
            write_literal_cell(_merged_anchor_cell(worksheet, row, note_column), descriptions)
            for issue in item.issues:
                issue_rows[issue.id] = row
            source_height = template_worksheet.row_dimensions[row].height or 27.0
            worksheet.row_dimensions[row].height = max(
                source_height,
                _template_height(item.content, chars_per_line=42, minimum=source_height),
                _template_height(descriptions, chars_per_line=24, minimum=source_height),
            )

    worksheet.print_area = template_worksheet.print_area
    worksheet.freeze_panes = template_worksheet.freeze_panes
    return issue_rows


def _build_audit_report_generic(
    worksheet,
    snapshot: ExportSnapshot,
    phase: ExportPhaseSnapshot,
    audit: ExportAuditSnapshot,
    checklist_sheet: str,
    checklist_issue_rows: dict[int, int],
) -> tuple[int, ...]:
    worksheet.sheet_view.showGridLines = False
    widths = (3, 10, 24, 12, 8, 12, 12, 14, 16, 24, 8)
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width
    _merge_write(worksheet, "B1:K1", "过程/产品审核报告", fill=_TITLE_FILL, bold=True)
    worksheet["B1"].font = Font(name="宋体", size=16, bold=True, color="FFFFFF")
    _merge_write(
        worksheet,
        "B2:K2",
        f"编号：ZDPX-{snapshot.project_code}/{snapshot.config_item_number}-QA-审核报告-{audit.sequence_no:03d}",
    )
    _merge_write(worksheet, "B3:C3", "项目名称", fill=_HEADER_FILL, bold=True)
    _merge_write(worksheet, "D3:G3", snapshot.project_name)
    _merge_write(worksheet, "H3:I3", "所处阶段", fill=_HEADER_FILL, bold=True)
    _merge_write(worksheet, "J3:K3", phase.name)
    _merge_write(worksheet, "B4:E4", "审核对象（过程名称/产品名称）", fill=_HEADER_FILL, bold=True)
    _merge_write(worksheet, "F4:K4", _audit_object(audit))
    _merge_write(worksheet, "B5:C5", "审核人", fill=_HEADER_FILL, bold=True)
    _merge_write(worksheet, "D5:E5", audit.auditor or snapshot.qa_person)
    _merge_write(worksheet, "F5:G5", "审核日期", fill=_HEADER_FILL, bold=True)
    _merge_write(worksheet, "H5:K5", audit.audit_date)
    _merge_write(worksheet, "B6:C6", "审核依据", fill=_HEADER_FILL, bold=True)
    _merge_write(
        worksheet,
        "D6:K6",
        f"ZDPX-{snapshot.project_code}/{snapshot.config_item_number}-QA-检查单-{audit.sequence_no:03d}",
    )
    issues = _all_issues(audit)
    overview = (
        f"本次审核依据《{audit.checklist_name}》实施。"
        + (f"经审核，发现{len(issues)}个不符合项。" if issues else "经审核，未发现不符合项。")
    )
    _merge_write(worksheet, "B7:K8", overview)
    summary_headers = ("NC总数", "", "轻微", "", "一般", "", "严重", "", "")
    for column, value in enumerate(summary_headers, start=2):
        worksheet.cell(9, column, value)
    _style_range(worksheet, "B9:K9", fill=_SUMMARY_FILL, bold=True)
    detail_start = 13
    detail_end = max(detail_start, detail_start + len(issues) - 1)
    if issues:
        _write_formula(worksheet["B10"], f'=COUNTIF(B{detail_start}:B{detail_end},"<>")')
        _write_formula(worksheet["D10"], f'=COUNTIF(F{detail_start}:F{detail_end},"轻微")')
        _write_formula(worksheet["F10"], f'=COUNTIF(F{detail_start}:F{detail_end},"一般")')
        _write_formula(worksheet["H10"], f'=COUNTIF(F{detail_start}:F{detail_end},"严重")')
    else:
        for coordinate in ("B10", "D10", "F10", "H10"):
            _write_formula(worksheet[coordinate], "=0")
    _style_range(worksheet, "B10:K10")
    _merge_write(worksheet, "B11:K11", "不符合项问题（NC）列表", fill=_HEADER_FILL, bold=True)
    headers = ("序号", "问题描述", "", "", "严重程度", "责任人", "提交日期", "计划关闭日期", "纠正措施", "")
    for column, value in enumerate(headers, start=2):
        worksheet.cell(12, column, value)
    _style_range(worksheet, "B12:K12", fill=_HEADER_FILL, bold=True)
    report_issue_rows = []
    for offset, issue in enumerate(issues):
        row = detail_start + offset
        source_row = checklist_issue_rows[issue.id]
        report_issue_rows.append(row)
        worksheet.cell(row, 2, offset + 1)
        _merge_write(worksheet, f"C{row}:E{row}", None)
        _write_formula(worksheet.cell(row, 3), f"={internal_reference(checklist_sheet, f'C{source_row}')}")
        for target_column, source_column in ((6, "F"), (7, "G"), (8, "H"), (9, "I")):
            _write_formula(
                worksheet.cell(row, target_column),
                f"={internal_reference(checklist_sheet, f'{source_column}{source_row}')}",
            )
        _merge_write(worksheet, f"J{row}:K{row}", None)
        _write_formula(worksheet.cell(row, 10), f"={internal_reference(checklist_sheet, f'J{source_row}')}")
        _style_range(worksheet, f"B{row}:B{row}")
        _style_range(worksheet, f"F{row}:I{row}")
        worksheet.row_dimensions[row].height = max(
            _wrapped_height(issue.description, chars_per_line=30, max_lines=3),
            _wrapped_height(issue.corrective_action, chars_per_line=18, max_lines=3),
        )
    if not issues:
        _merge_write(worksheet, f"B{detail_start}:K{detail_start}", "本次审核未发现不符合项。")
    verification_row = detail_end + 2
    _merge_write(worksheet, f"B{verification_row}:K{verification_row}", "纠正措施验证情况：", fill=_HEADER_FILL, bold=True)
    verification = "；".join(
        (
            f"待确认的建议值：{issue.verification_conclusion}"
            if issue.suggested_closed and not issue.confirmed and issue.verification_conclusion
            else issue.verification_conclusion or "尚未完成验证"
        )
        for issue in issues
    ) or "不适用"
    _merge_write(worksheet, f"B{verification_row + 1}:K{verification_row + 2}", verification)
    worksheet.freeze_panes = "B12"
    worksheet.print_title_rows = "1:12"
    worksheet.print_area = f"B1:K{verification_row + 2}"
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_setup.paperSize = worksheet.PAPERSIZE_A4
    worksheet.page_setup.orientation = worksheet.ORIENTATION_LANDSCAPE
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.print_options.horizontalCentered = True
    return tuple(report_issue_rows)


def _build_audit_report(
    worksheet,
    snapshot: ExportSnapshot,
    phase: ExportPhaseSnapshot,
    audit: ExportAuditSnapshot,
    checklist_sheet: str,
    checklist_issue_rows: dict[int, int],
    template_worksheet=None,
) -> tuple[int, ...]:
    """Build a report using the authoritative audit-report layout."""

    if template_worksheet is None:
        return _build_audit_report_generic(
            worksheet,
            snapshot,
            phase,
            audit,
            checklist_sheet,
            checklist_issue_rows,
        )

    original_title = worksheet.title
    copy_worksheet_template(template_worksheet, worksheet)
    worksheet.title = original_title
    issues = _all_issues(audit)
    detail_rows = max(2, len(issues))

    for merged_range in tuple(worksheet.merged_cells.ranges):
        if merged_range.min_row >= 13:
            worksheet.unmerge_cells(str(merged_range))
    if detail_rows > 2:
        worksheet.insert_rows(15, detail_rows - 2)
    _clear_row_values(worksheet, 1, max(worksheet.max_row, 18), max_column=11)

    for offset in range(detail_rows):
        row = 13 + offset
        _copy_row_format(template_worksheet, 13, worksheet, row, 11)
        worksheet.merge_cells(start_row=row, start_column=3, end_row=row, end_column=5)
        worksheet.merge_cells(start_row=row, start_column=10, end_row=row, end_column=11)
        worksheet.row_dimensions[row].height = 72.0

    verification_header_row = 13 + detail_rows
    verification_body_row = verification_header_row + 1
    _copy_row_format(template_worksheet, 15, worksheet, verification_header_row, 11)
    _copy_row_format(template_worksheet, 16, worksheet, verification_body_row, 11)
    worksheet.merge_cells(
        start_row=verification_header_row,
        start_column=2,
        end_row=verification_header_row,
        end_column=11,
    )
    worksheet.merge_cells(
        start_row=verification_body_row,
        start_column=2,
        end_row=verification_body_row,
        end_column=11,
    )

    worksheet["B1"] = "过程/产品审核报告"
    worksheet["B2"] = (
        f"编号：ZDPX-{snapshot.project_code}/{snapshot.config_item_number}"
        f"-QA-审核报告-{audit.sequence_no:03d}"
    )
    worksheet["B3"] = "项目名称"
    write_literal_cell(worksheet["C3"], snapshot.project_name)
    worksheet["H3"] = "所处阶段"
    write_literal_cell(worksheet["I3"], phase.name)
    worksheet["B4"] = "审核对象（过程名称/产品名称）"
    write_literal_cell(worksheet["F4"], _audit_object(audit))
    worksheet["B5"] = "审核人"
    write_literal_cell(worksheet["C5"], audit.auditor or snapshot.qa_person)
    worksheet["F5"] = "审核日期"
    write_literal_cell(worksheet["H5"], audit.audit_date)
    worksheet["B6"] = "审核依据"
    worksheet["C6"] = (
        f"ZDPX-{snapshot.project_code}/{snapshot.config_item_number}"
        f"-QA-检查单-{audit.sequence_no:03d}"
    )
    worksheet["B7"] = "审核内容概述："
    overview = _audit_overview(audit)
    write_literal_cell(worksheet["B8"], overview)
    worksheet.row_dimensions[8].height = _template_height(
        overview,
        chars_per_line=48,
        minimum=128.5,
        line_height=18.0,
    )

    worksheet["B9"] = "按严重程度统计NC"
    worksheet["D9"] = "NC总数"
    worksheet["F9"] = "轻 微"
    worksheet["H9"] = "一 般"
    worksheet["J9"] = "严 重"
    detail_end = 12 + detail_rows
    _write_formula(worksheet["D10"], f'=COUNTIF(B13:B{detail_end},"?*")')
    _write_formula(worksheet["F10"], f'=COUNTIF(F13:F{detail_end},"轻微")')
    _write_formula(worksheet["H10"], f'=COUNTIF(F13:F{detail_end},"一般")')
    _write_formula(worksheet["J10"], f'=COUNTIF(F13:F{detail_end},"严重")')
    worksheet["B11"] = "不符合项问题（NC）列表"
    for coordinate, value in (
        ("B12", "序 号"),
        ("C12", "问题描述"),
        ("F12", "严重程度"),
        ("G12", "责任人"),
        ("H12", "提交日期"),
        ("I12", "计划关闭日期"),
        ("J12", "纠正措施"),
    ):
        worksheet[coordinate] = value

    report_issue_rows = []
    for offset, issue in enumerate(issues):
        row = 13 + offset
        report_issue_rows.append(row)
        worksheet.cell(row, 2, offset + 1)
        write_literal_cell(worksheet.cell(row, 3), issue.description)
        write_literal_cell(worksheet.cell(row, 6), issue.severity)
        write_literal_cell(worksheet.cell(row, 7), issue.responsible_person)
        write_literal_cell(worksheet.cell(row, 8), issue.submitted_date)
        write_literal_cell(worksheet.cell(row, 9), issue.planned_close_date)
        corrective_action = issue.corrective_action
        if issue.suggested_closed and not issue.confirmed and corrective_action:
            corrective_action = f"待确认的建议值：{corrective_action}"
        write_literal_cell(worksheet.cell(row, 10), corrective_action)
        worksheet.row_dimensions[row].height = max(
            72.0,
            _template_height(issue.description, chars_per_line=26, minimum=72.0),
            _template_height(corrective_action, chars_per_line=24, minimum=72.0),
        )

    worksheet.cell(verification_header_row, 2, "纠正措施验证情况：")
    conclusions = list(
        dict.fromkeys(
            (
                f"待确认的建议值：{issue.verification_conclusion}"
                if issue.suggested_closed and not issue.confirmed and issue.verification_conclusion
                else issue.verification_conclusion
            )
            for issue in issues
            if issue.verification_conclusion
        )
    )
    conclusion_text = "；".join(conclusions)
    if issues and not conclusion_text:
        conclusion_text = "尚未完成验证。"
    verification = f"{conclusion_text}\n\n                                                     验证人/日期："
    worksheet.cell(verification_body_row, 2, verification)
    worksheet.row_dimensions[verification_body_row].height = 94.25 if conclusion_text else 72.0

    _apply_outer_border(
        worksheet,
        min_row=3,
        min_column=2,
        max_row=verification_body_row,
        max_column=11,
        side=_THICK,
    )

    worksheet.print_area = f"B1:K{verification_body_row}"
    worksheet.page_setup.orientation = worksheet.ORIENTATION_PORTRAIT
    worksheet.page_setup.fitToWidth = None
    worksheet.page_setup.fitToHeight = None
    worksheet.freeze_panes = None
    worksheet.print_title_rows = None
    return tuple(report_issue_rows)


def trend_summary(phase: ExportPhaseSnapshot) -> dict[str, int]:
    counts = dict.fromkeys(("NC总数", "轻微", "一般", "严重", *_TREND_CATEGORIES), 0)
    for audit in phase.audits:
        if audit.report_scope != "phase":
            continue
        for issue in _all_issues(audit):
            counts["NC总数"] += 1
            counts[issue.severity] += 1
            category = "工作产品" if audit.audit_kind == "product" else (audit.process_code or "过程")
            counts[category] = counts.get(category, 0) + 1
    return counts


def _build_calculation_sheet(
    worksheet, infos: tuple[_AuditSheetInfo, ...]
) -> tuple[int, dict[tuple[int, int], int]]:
    headers = (
        "流水号", "审核对象", "问题序号", "NC类型", "问题描述", "严重程度",
        "是否关闭", "归组日期", "节点类型", "审核日期", "节点ID",
    )
    for column, value in enumerate(headers, start=1):
        worksheet.cell(1, column, value)
    row = 2
    rows_by_issue: dict[tuple[int, int], int] = {}
    for info in infos:
        audit = info.audit
        issues = _all_issues(audit)
        for index, report_row in enumerate(info.report_issue_rows, start=1):
            issue = issues[index - 1]
            rows_by_issue[(audit.id, issue.id)] = row
            worksheet.cell(row, 1, audit.sequence_no)
            _write_formula(worksheet.cell(row, 2), f"={internal_reference(info.report_sheet, 'F4')}")
            _write_formula(worksheet.cell(row, 3), f"={internal_reference(info.report_sheet, f'B{report_row}')}")
            write_literal_cell(worksheet.cell(row, 4), "工作产品" if audit.audit_kind == "product" else (audit.process_code or "过程"))
            _write_formula(worksheet.cell(row, 5), f"={internal_reference(info.report_sheet, f'C{report_row}')}")
            _write_formula(worksheet.cell(row, 6), f"={internal_reference(info.report_sheet, f'F{report_row}')}")
            worksheet.cell(row, 7, "是" if issue.confirmed else "否")
            worksheet.cell(row, 8, audit.tracking_node_date or "未匹配")
            worksheet.cell(row, 9, audit.tracking_node_type or "未匹配节点")
            worksheet.cell(row, 10, audit.audit_date)
            worksheet.cell(row, 11, audit.tracking_node_id)
            row += 1
    issue_last_row = row - 1
    calc_end = max(2, issue_last_row)
    type_range = f"D2:D{calc_end}"
    severity_range = f"F2:F{calc_end}"
    summary_row = row + 1
    summary_formulas = (
        ("NC总数", f'=COUNTIF({type_range},"?*")'),
        ("轻微", f'=COUNTIF({severity_range},"轻微")'),
        ("一般", f'=COUNTIF({severity_range},"一般")'),
        ("严重", f'=COUNTIF({severity_range},"严重")'),
        *( 
            (category, f'=COUNTIF({type_range},"{category}")')
            for category in _TREND_CATEGORIES
        ),
    )
    for offset, (label, expression) in enumerate(summary_formulas):
        worksheet.cell(summary_row + offset, 1, label)
        _write_formula(worksheet.cell(summary_row + offset, 2), expression)
    worksheet.sheet_state = "hidden"
    return issue_last_row, rows_by_issue


def _build_phase_report_generic(
    worksheet,
    snapshot: ExportSnapshot,
    phase: ExportPhaseSnapshot,
    infos: tuple[_AuditSheetInfo, ...],
    calculation_last_row: int,
    calculation_rows: dict[tuple[int, int], int],
) -> None:
    worksheet.sheet_view.showGridLines = False
    widths = (10, 22, 12, 12, 12, 12, 12, 12, 12, 14, 12, 12)
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width
    _merge_write(worksheet, "A1:L1", "质量审核工作报告", fill=_TITLE_FILL, bold=True)
    worksheet["A1"].font = Font(name="宋体", size=18, bold=True, color="FFFFFF")
    _merge_write(
        worksheet,
        "A2:L2",
        f"编号：ZDPX-{snapshot.project_code}/{snapshot.config_item_number}-QA-质量报告-{phase.sort_order:03d}",
    )
    _merge_write(worksheet, "A3:B3", "项目名称", fill=_HEADER_FILL, bold=True)
    _merge_write(worksheet, "C3:F3", snapshot.project_name)
    _merge_write(worksheet, "G3:H3", "报告阶段", fill=_HEADER_FILL, bold=True)
    _merge_write(worksheet, "I3:L3", phase.name)
    _merge_write(worksheet, "A4:B4", "报告起止日期", fill=_HEADER_FILL, bold=True)
    _merge_write(worksheet, "C4:F4", f"{phase.start_date} 至 {phase.end_date}")
    _merge_write(worksheet, "G4:H4", "项目QA", fill=_HEADER_FILL, bold=True)
    _merge_write(worksheet, "I4:L4", snapshot.qa_person)

    _merge_write(worksheet, "A6:L6", "审核情况汇总", fill=_HEADER_FILL, bold=True)
    headers = ("流水号", "审核对象", "审核日期", "节点", "NC类型", "问题描述", "", "", "", "严重程度", "是否关闭", "审核状态")
    for column, value in enumerate(headers, start=1):
        worksheet.cell(7, column, value)
    _style_range(worksheet, "A7:L7", fill=_HEADER_FILL, bold=True)
    row = 8
    node_keys = [
        (node.id, node.node_date, node.node_type) for node in phase.tracking_nodes
    ]
    if any(audit.tracking_node_date is None for audit in (info.audit for info in infos)):
        node_keys.append((None, "未匹配", "未匹配节点"))
    if not node_keys:
        node_keys.append((None, phase.end_date, "阶段"))
    for node_id, node_date, node_type in node_keys:
        matching = [
            info for info in infos
            if info.audit.tracking_node_id == node_id
        ]
        _merge_write(worksheet, f"A{row}:L{row}", f"{node_type}汇总：{node_date}", fill=_GROUP_FILL, bold=True)
        row += 1
        block_start = row
        for info in matching:
            issues = _all_issues(info.audit)
            if not issues:
                worksheet.cell(row, 1, info.audit.sequence_no)
                write_literal_cell(worksheet.cell(row, 2), _audit_object(info.audit))
                worksheet.cell(row, 3, info.audit.audit_date)
                worksheet.cell(row, 4, node_type)
                worksheet.cell(row, 12, _audit_state_label(info.audit.confirmation_state))
                _merge_write(worksheet, f"F{row}:I{row}", "未发现不符合项")
                _style_range(worksheet, f"A{row}:E{row}")
                _style_range(worksheet, f"J{row}:L{row}")
                worksheet.row_dimensions[row].height = _wrapped_height(
                    _audit_object(info.audit), chars_per_line=12, max_lines=2
                )
                row += 1
                continue
            for issue in issues:
                calc_row = calculation_rows[(info.audit.id, issue.id)]
                worksheet.cell(row, 1, info.audit.sequence_no)
                _write_formula(worksheet.cell(row, 2), f"={internal_reference(CALCULATION_SHEET, f'B{calc_row}')}")
                worksheet.cell(row, 3, info.audit.audit_date)
                worksheet.cell(row, 4, node_type)
                for target_column, source_column in ((5, "D"), (6, "E"), (10, "F"), (11, "G")):
                    _write_formula(worksheet.cell(row, target_column), f"={internal_reference(CALCULATION_SHEET, f'{source_column}{calc_row}')}")
                worksheet.merge_cells(start_row=row, start_column=6, end_row=row, end_column=9)
                worksheet.cell(row, 12, _audit_state_label(info.audit.confirmation_state))
                _style_range(worksheet, f"A{row}:L{row}")
                worksheet.row_dimensions[row].height = max(
                    _wrapped_height(_audit_object(info.audit), chars_per_line=12, max_lines=2),
                    _wrapped_height(issue.description, chars_per_line=32, max_lines=2),
                )
                row += 1
        if not matching:
            _merge_write(worksheet, f"A{row}:L{row}", "该节点无审核记录。")
            row += 1
        block_end = row - 1
        _merge_write(worksheet, f"A{row}:D{row}", "节点统计", fill=_SUMMARY_FILL, bold=True)
        worksheet.cell(row, 5, "NC总数")
        _write_formula(worksheet.cell(row, 6), f'=COUNTIF(J{block_start}:J{block_end},"?*")')
        worksheet.cell(row, 7, "未关闭")
        _write_formula(worksheet.cell(row, 8), f'=COUNTIF(K{block_start}:K{block_end},"否")')
        worksheet.cell(row, 9, "严重")
        _write_formula(worksheet.cell(row, 10), f'=COUNTIF(J{block_start}:J{block_end},"严重")')
        _style_range(worksheet, f"E{row}:L{row}", fill=_SUMMARY_FILL, bold=True)
        row += 2

    _merge_write(worksheet, f"A{row}:L{row}", "阶段汇总统计", fill=_TITLE_FILL, bold=True)
    worksheet[f"A{row}"].font = Font(name="宋体", size=12, bold=True, color="FFFFFF")
    row += 1
    calc_end = max(2, calculation_last_row)
    type_range = internal_reference(CALCULATION_SHEET, f"D2:D{calc_end}")
    closed_range = internal_reference(CALCULATION_SHEET, f"G2:G{calc_end}")
    severity_range = internal_reference(CALCULATION_SHEET, f"F2:F{calc_end}")
    metrics = (
        ("产品NC总数", f'=COUNTIF({type_range},"工作产品")'),
        ("过程NC总数", f'=COUNTIFS({type_range},"<>工作产品",{type_range},"<>")'),
        ("未关闭NC数", f'=COUNTIF({closed_range},"否")'),
        ("轻微", f'=COUNTIF({severity_range},"轻微")'),
        ("一般", f'=COUNTIF({severity_range},"一般")'),
        ("严重", f'=COUNTIF({severity_range},"严重")'),
    )
    for index, (label, expression) in enumerate(metrics):
        column = 1 + index * 2
        worksheet.cell(row, column, label)
        _write_formula(worksheet.cell(row, column + 1), expression)
    _style_range(worksheet, f"A{row}:L{row}", fill=_SUMMARY_FILL, bold=True)
    row += 1
    # Keep this sequence aligned with the Step 9 trend workbook.  Every
    # category must be present even when its current count is zero, otherwise
    # a cross-workbook MATCH() would silently turn a real category into zero.
    codes = ("RDM", "PP", "PMC", "CM", "MT", "MPM", "VV", "LD", "II", "工作产品")
    for column, code in enumerate(codes, start=1):
        worksheet.cell(row, column, code)
        _write_formula(
            worksheet.cell(row + 1, column),
            f'=COUNTIF({type_range},"{code}")',
        )
    worksheet.cell(row, 11, "总计")
    _write_formula(worksheet.cell(row + 1, 11), f"=SUM(A{row + 1}:J{row + 1})")
    _style_range(worksheet, f"A{row}:K{row + 1}", fill=_HEADER_FILL, bold=True)
    worksheet.freeze_panes = "A8"
    worksheet.print_title_rows = "1:7"
    worksheet.print_area = f"A1:L{row + 1}"
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_setup.paperSize = worksheet.PAPERSIZE_A4
    worksheet.page_setup.orientation = worksheet.ORIENTATION_LANDSCAPE
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 1
    worksheet.print_options.horizontalCentered = True


def _countifs_formula(pairs: tuple[tuple[str, str | int], ...]) -> str:
    arguments = []
    for cell_range, criterion in pairs:
        arguments.append(cell_range)
        arguments.append(str(criterion) if isinstance(criterion, int) else f'"{criterion}"')
    return f"=COUNTIFS({','.join(arguments)})"


def _populate_summary_block(
    worksheet,
    template_worksheet,
    target_row: int,
    *,
    template_row: int,
    period_label: str,
    audits: tuple[ExportAuditSnapshot, ...],
    start_date: str,
    end_date: str,
    calculation_last_row: int,
    node_id: int | None,
) -> int:
    copy_worksheet_rows(
        template_worksheet,
        worksheet,
        template_row,
        template_row + 17,
        target_row,
    )
    _clear_row_values(worksheet, target_row, target_row + 17)

    calc_end = max(2, calculation_last_row)
    type_range = internal_reference(CALCULATION_SHEET, f"$D$2:$D${calc_end}")
    severity_range = internal_reference(CALCULATION_SHEET, f"$F$2:$F${calc_end}")
    closed_range = internal_reference(CALCULATION_SHEET, f"$G$2:$G${calc_end}")
    node_range = internal_reference(CALCULATION_SHEET, f"$K$2:$K${calc_end}")

    def pairs(*criteria: tuple[str, str | int]) -> tuple[tuple[str, str | int], ...]:
        result = list(criteria)
        if node_id is not None:
            result.append((node_range, node_id))
        return tuple(result)

    worksheet.cell(target_row, 1, f"{period_label}汇总统计：")
    worksheet.cell(target_row, 3, _audit_scope_text(audits, period_label))
    worksheet.row_dimensions[target_row].height = _template_height(
        worksheet.cell(target_row, 3).value,
        chars_per_line=60,
        minimum=template_worksheet.row_dimensions[template_row].height or 59.5,
        line_height=18.0,
    )
    worksheet.cell(target_row + 1, 1, "起始时间")
    worksheet.cell(target_row + 1, 3, start_date)
    worksheet.cell(target_row + 1, 7, "结束时间")
    worksheet.cell(target_row + 1, 9, end_date)

    product_total_row = target_row + 4
    product_value_row = target_row + 5
    product_ratio_row = target_row + 6
    worksheet.cell(target_row + 2, 1, "产品NC统计")
    worksheet.cell(target_row + 3, 1, "产品NC总数")
    worksheet.cell(target_row + 3, 3, "关闭情况")
    worksheet.cell(target_row + 3, 7, "按严重程度分布情况")
    _write_formula(
        worksheet.cell(product_total_row, 1),
        _countifs_formula(pairs((type_range, "工作产品"))),
    )
    worksheet.cell(product_total_row, 3, "未关闭的产品NC数")
    worksheet.cell(product_total_row, 5, "产品NC关闭率")
    worksheet.cell(product_total_row, 9, "轻微")
    worksheet.cell(product_total_row, 11, "一般")
    worksheet.cell(product_total_row, 12, "严重")
    _write_formula(
        worksheet.cell(product_value_row, 3),
        _countifs_formula(pairs((type_range, "工作产品"), (closed_range, "否"))),
    )
    _write_formula(
        worksheet.cell(product_value_row, 5),
        f'=IF(A{product_total_row}=0,"—",(A{product_total_row}-C{product_value_row})/A{product_total_row})',
    )
    worksheet.cell(product_value_row, 7, "NC数")
    for column, severity in ((9, "轻微"), (11, "一般"), (12, "严重")):
        _write_formula(
            worksheet.cell(product_value_row, column),
            _countifs_formula(
                pairs((type_range, "工作产品"), (severity_range, severity))
            ),
        )
    worksheet.cell(product_ratio_row, 7, "所占比例")
    for column in (9, 11, 12):
        coordinate = worksheet.cell(product_value_row, column).coordinate
        _write_formula(
            worksheet.cell(product_ratio_row, column),
            f'=IF(A{product_total_row}=0,"—",{coordinate}/A{product_total_row})',
        )
        worksheet.cell(product_ratio_row, column).number_format = "0.00%"
    worksheet.cell(product_value_row, 5).number_format = "0.00%"

    process_total_row = target_row + 9
    process_value_row = target_row + 10
    process_ratio_row = target_row + 11
    worksheet.cell(target_row + 7, 1, "过程NC统计")
    worksheet.cell(target_row + 8, 1, "过程NC总数")
    worksheet.cell(target_row + 8, 3, "关闭情况")
    worksheet.cell(target_row + 8, 7, "按严重程度分布情况")
    _write_formula(
        worksheet.cell(process_total_row, 1),
        _countifs_formula(
            pairs((type_range, "<>工作产品"), (type_range, "<>"))
        ),
    )
    worksheet.cell(process_total_row, 3, "未关闭的过程NC数")
    worksheet.cell(process_total_row, 5, "过程NC关闭率")
    worksheet.cell(process_total_row, 9, "轻微")
    worksheet.cell(process_total_row, 11, "一般")
    worksheet.cell(process_total_row, 12, "严重")
    _write_formula(
        worksheet.cell(process_value_row, 3),
        _countifs_formula(
            pairs(
                (type_range, "<>工作产品"),
                (type_range, "<>"),
                (closed_range, "否"),
            )
        ),
    )
    _write_formula(
        worksheet.cell(process_value_row, 5),
        f'=IF(A{process_total_row}=0,"—",(A{process_total_row}-C{process_value_row})/A{process_total_row})',
    )
    worksheet.cell(process_value_row, 7, "NC数")
    for column, severity in ((9, "轻微"), (11, "一般"), (12, "严重")):
        _write_formula(
            worksheet.cell(process_value_row, column),
            _countifs_formula(
                pairs(
                    (type_range, "<>工作产品"),
                    (type_range, "<>"),
                    (severity_range, severity),
                )
            ),
        )
    worksheet.cell(process_ratio_row, 7, "所占比例")
    for column in (9, 11, 12):
        coordinate = worksheet.cell(process_value_row, column).coordinate
        _write_formula(
            worksheet.cell(process_ratio_row, column),
            f'=IF(A{process_total_row}=0,"—",{coordinate}/A{process_total_row})',
        )
        worksheet.cell(process_ratio_row, column).number_format = "0.00%"
    worksheet.cell(process_value_row, 5).number_format = "0.00%"

    distribution_header_row = target_row + 13
    distribution_value_row = target_row + 14
    distribution_ratio_row = target_row + 15
    worksheet.cell(target_row + 12, 1, "过程及产品NC占比统计")
    worksheet.cell(distribution_header_row, 1, "NC分布情况")
    for column, category in enumerate(_PHASE_VISIBLE_CATEGORIES, start=3):
        worksheet.cell(distribution_header_row, column, category)
        _write_formula(
            worksheet.cell(distribution_value_row, column),
            _countifs_formula(pairs((type_range, category))),
        )
    worksheet.cell(distribution_header_row, 12, "总计")
    worksheet.cell(distribution_value_row, 1, "NC数")
    _write_formula(
        worksheet.cell(distribution_value_row, 12),
        f"=SUM(C{distribution_value_row}:K{distribution_value_row})",
    )
    worksheet.cell(distribution_ratio_row, 1, "所占比例")
    for column in range(3, 12):
        coordinate = worksheet.cell(distribution_value_row, column).coordinate
        _write_formula(
            worksheet.cell(distribution_ratio_row, column),
            f'=IF(L{distribution_value_row}=0,"—",{coordinate}/L{distribution_value_row})',
        )
        worksheet.cell(distribution_ratio_row, column).number_format = "0.00%"
    worksheet.cell(distribution_ratio_row, 12, "─")
    worksheet.cell(target_row + 16, 1, "数据分析")
    analysis = _analysis_text(audits, period_label)
    worksheet.cell(target_row + 17, 1, analysis)
    worksheet.row_dimensions[target_row + 17].height = _template_height(
        analysis,
        chars_per_line=70,
        minimum=template_worksheet.row_dimensions[template_row + 17].height or 54.0,
        line_height=18.0,
    )
    return target_row + 18


def _build_phase_report(
    worksheet,
    snapshot: ExportSnapshot,
    phase: ExportPhaseSnapshot,
    infos: tuple[_AuditSheetInfo, ...],
    calculation_last_row: int,
    calculation_rows: dict[tuple[int, int], int],
    template_worksheet=None,
) -> None:
    if template_worksheet is None:
        _build_phase_report_generic(
            worksheet,
            snapshot,
            phase,
            infos,
            calculation_last_row,
            calculation_rows,
        )
        return

    original_title = worksheet.title
    copy_worksheet_columns(template_worksheet, worksheet)
    copy_worksheet_rows(template_worksheet, worksheet, 1, 5, 1)
    worksheet.title = original_title
    _clear_row_values(worksheet, 1, 5)
    worksheet["A1"] = "质量审核工作报告"
    worksheet["A2"] = (
        f"编号：ZDPX-{snapshot.project_code}/{snapshot.config_item_number}"
        f"-QA-质量报告-{phase.sort_order:03d}"
    )
    worksheet["A3"] = "项目名称"
    write_literal_cell(worksheet["B3"], snapshot.project_name)
    worksheet["A4"] = "报告阶段"
    write_literal_cell(worksheet["B4"], phase.name)
    worksheet["D4"] = "报告起止日期"
    write_literal_cell(worksheet["F4"], f"{phase.start_date} 至 {phase.end_date}")
    worksheet["J4"] = "项目QA"
    write_literal_cell(worksheet["K4"], snapshot.qa_person)
    worksheet["A5"] = "审核情况汇总"

    row = 6

    def append_issue_blocks(block_infos: tuple[_AuditSheetInfo, ...]) -> None:
        nonlocal row
        for info in block_infos:
            issues = _all_issues(info.audit)
            if not issues:
                continue
            copy_worksheet_rows(template_worksheet, worksheet, 6, 7, row)
            _clear_row_values(worksheet, row, row + 1)
            worksheet.cell(row, 1, "审核过程/产品名称")
            _write_formula(
                worksheet.cell(row, 3),
                f"={internal_reference(info.report_sheet, 'F4')}",
            )
            worksheet.cell(row, 11, "审核时间")
            _write_formula(
                worksheet.cell(row, 12),
                f"={internal_reference(info.report_sheet, 'H5')}",
            )
            worksheet.cell(row + 1, 1, "不符合项列表")
            worksheet.cell(row + 1, 10, "NC类型")
            worksheet.cell(row + 1, 11, "严重程度")
            worksheet.cell(row + 1, 12, "是否关闭")
            row += 2
            for issue, report_row in zip(issues, info.report_issue_rows, strict=True):
                copy_worksheet_rows(template_worksheet, worksheet, 8, 8, row)
                _clear_row_values(worksheet, row, row)
                _write_formula(
                    worksheet.cell(row, 1),
                    f"={internal_reference(info.report_sheet, f'B{report_row}')}",
                )
                _write_formula(
                    worksheet.cell(row, 2),
                    f"={internal_reference(info.report_sheet, f'C{report_row}')}",
                )
                worksheet.cell(
                    row,
                    10,
                    "工作产品" if info.audit.audit_kind == "product" else (info.audit.process_code or "过程"),
                )
                _write_formula(
                    worksheet.cell(row, 11),
                    f"={internal_reference(info.report_sheet, f'F{report_row}')}",
                )
                worksheet.cell(row, 12, "是" if issue.confirmed else "否")
                worksheet.row_dimensions[row].height = _template_height(
                    issue.description,
                    chars_per_line=58,
                    minimum=template_worksheet.row_dimensions[8].height or 19.75,
                    line_height=18.0,
                )
                row += 1

    non_stage_nodes = tuple(node for node in phase.tracking_nodes if node.node_type != "阶段")
    consumed_ids: set[int] = set()
    previous_end = _iso_date(phase.start_date) - timedelta(days=1)
    for node in non_stage_nodes:
        matching = tuple(info for info in infos if info.audit.tracking_node_id == node.id)
        consumed_ids.update(info.audit.id for info in matching)
        append_issue_blocks(matching)
        period_start = previous_end + timedelta(days=1)
        period_end = _iso_date(node.node_date)
        row = _populate_summary_block(
            worksheet,
            template_worksheet,
            row,
            template_row=22,
            period_label=node.node_type,
            audits=tuple(info.audit for info in matching),
            start_date=period_start.isoformat(),
            end_date=period_end.isoformat(),
            calculation_last_row=calculation_last_row,
            node_id=node.id,
        )
        previous_end = period_end

    remaining = tuple(info for info in infos if info.audit.id not in consumed_ids)
    append_issue_blocks(remaining)
    row = _populate_summary_block(
        worksheet,
        template_worksheet,
        row,
        template_row=87,
        period_label="阶段",
        audits=tuple(info.audit for info in infos),
        start_date=phase.start_date,
        end_date=phase.end_date,
        calculation_last_row=calculation_last_row,
        node_id=None,
    )

    worksheet.print_area = f"A1:L{row - 1}"
    worksheet.page_setup.orientation = worksheet.ORIENTATION_LANDSCAPE
    worksheet.page_setup.fitToHeight = 0
    worksheet.freeze_panes = None
    worksheet.print_title_rows = None


def build_phase_workbook(
    snapshot: ExportSnapshot,
    phase: ExportPhaseSnapshot,
    output_path: Path,
) -> PhaseWorkbookArtifact:
    if phase not in snapshot.phases:
        raise ValueError("Phase does not belong to the export snapshot")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    _set_recalculation(workbook)
    product_templates = load_workbook(_PRODUCT_CHECKLIST_REFERENCE, data_only=False)
    process_templates = load_workbook(_PROCESS_CHECKLIST_REFERENCE, data_only=False)
    report_templates = load_workbook(_REPORT_REFERENCE, data_only=False)
    try:
        template_books = {
            _PRODUCT_CHECKLIST_REFERENCE.name: product_templates,
            _PROCESS_CHECKLIST_REFERENCE.name: process_templates,
        }
        used = {DIRECTORY_SHEET, PHASE_REPORT_SHEET, CALCULATION_SHEET}
        phase_audits = tuple(audit for audit in phase.audits if audit.report_scope == "phase")
        provisional = []
        for audit in phase_audits:
            short = _audit_object(audit)
            checklist_sheet = safe_sheet_name(short, prefix=f"C{audit.sequence_no:03d}-", used=used)
            report_sheet = safe_sheet_name(short, prefix=f"R{audit.sequence_no:03d}-", used=used)
            source_book = template_books.get(audit.checklist_source_workbook or "")
            source_sheet = None
            if source_book is not None and audit.checklist_source_sheet in source_book.sheetnames:
                source_sheet = source_book[audit.checklist_source_sheet]
            checklist = workbook.create_sheet(checklist_sheet)
            issue_rows = _build_checklist(
                checklist,
                snapshot,
                phase,
                audit,
                source_sheet,
            )
            report = workbook.create_sheet(report_sheet)
            report_issue_rows = _build_audit_report(
                report,
                snapshot,
                phase,
                audit,
                checklist_sheet,
                issue_rows,
                report_templates[_AUDIT_REPORT_TEMPLATE_SHEET],
            )
            provisional.append(_AuditSheetInfo(audit, checklist_sheet, report_sheet, report_issue_rows))
        infos = tuple(provisional)
        calculation = workbook.create_sheet(CALCULATION_SHEET)
        calculation_last_row, calculation_rows = _build_calculation_sheet(calculation, infos)
        phase_report = workbook.create_sheet(PHASE_REPORT_SHEET, 1)
        _build_phase_report(
            phase_report,
            snapshot,
            phase,
            infos,
            calculation_last_row,
            calculation_rows,
            report_templates[_PHASE_REPORT_TEMPLATE_SHEET],
        )
        _build_directory(workbook, snapshot, phase, infos)
        workbook.save(output_path)
        counts = trend_summary(phase)
        summary_start = calculation_last_row + 2
        write_formula_caches(output_path, {CALCULATION_SHEET: {
            f"B{summary_start + offset}": counts[label]
            for offset, label in enumerate(("NC总数", "轻微", "一般", "严重", *_TREND_CATEGORIES))
        }})
    finally:
        workbook.close()
        product_templates.close()
        process_templates.close()
        report_templates.close()
    return PhaseWorkbookArtifact(
        phase_id=phase.id,
        phase_name=phase.name,
        path=output_path,
        audit_count=len(infos),
        issue_count=sum(len(_all_issues(info.audit)) for info in infos),
    )


def build_all_phase_workbooks(
    snapshot: ExportSnapshot,
    output_dir: Path,
) -> tuple[PhaseWorkbookArtifact, ...]:
    artifacts = []
    for phase in snapshot.phases:
        filename = safe_filename_component(
            f"QualityM-{snapshot.project_code}-V{snapshot.version_no}-{phase.sort_order:02d}-{phase.name}-阶段质量审核记录"
        ) + ".xlsx"
        artifacts.append(build_phase_workbook(snapshot, phase, output_dir / filename))
    return tuple(artifacts)
