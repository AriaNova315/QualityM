from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import LineChart, PieChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from app.application.export_snapshot_service import ExportSnapshot
from app.excel.external_links import add_external_workbook_link
from app.excel.formulas import Formula, external_reference, write_formula, write_literal_cell
from app.excel.phase_workbook_builder import (
    CALCULATION_SHEET, _AUDIT_REPORT_TEMPLATE_SHEET,
    _AuditSheetInfo, _PRODUCT_CHECKLIST_REFERENCE, _PROCESS_CHECKLIST_REFERENCE,
    _REPORT_REFERENCE, _audit_object, _build_audit_report, _build_checklist,
    _set_recalculation, _style_range,
)
from app.excel.sheet_names import safe_filename_component, safe_sheet_name

COMPANY_DIRECTORY_SHEET = "审核目录"
TREND_SHEET = "质量工作趋势分析"
_TITLE = PatternFill("solid", fgColor="1F4E78")
_HEADER = PatternFill("solid", fgColor="D9EAF7")
_THIN = Side(style="thin", color="000000")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_CATEGORIES = ("RDM", "PP", "PMC", "CM", "MT", "MPM", "VV", "LD", "II", "工作产品")

@dataclass(frozen=True, slots=True)
class CompanyTrendArtifacts:
    output_dir: Path
    phase_workbooks: tuple[Path, ...]
    company_workbook: Path | None
    trend_workbook: Path
    manifest: Path

def _style(ws, ref: str, fill=None) -> None:
    for row in ws[ref]:
        for cell in row:
            cell.border = _BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if fill: cell.fill = fill

def build_company_workbook(snapshot: ExportSnapshot, output_path: Path) -> Path | None:
    entries = [(phase, audit) for phase in snapshot.phases for audit in phase.audits if audit.report_scope == "company"]
    if not entries:
        return None
    output_path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook(); _set_recalculation(book)
    product_templates = load_workbook(_PRODUCT_CHECKLIST_REFERENCE, data_only=False)
    process_templates = load_workbook(_PROCESS_CHECKLIST_REFERENCE, data_only=False)
    report_templates = load_workbook(_REPORT_REFERENCE, data_only=False)
    try:
        template_books = {
            _PRODUCT_CHECKLIST_REFERENCE.name: product_templates,
            _PROCESS_CHECKLIST_REFERENCE.name: process_templates,
        }
        directory = book.active; directory.title = COMPANY_DIRECTORY_SHEET
        used = {COMPANY_DIRECTORY_SHEET}; infos = []
        for phase, audit in entries:
            base = f"{phase.sort_order:02d}-{_audit_object(audit)}"
            checklist_name = safe_sheet_name(base, prefix=f"C{audit.sequence_no:03d}-", used=used)
            report_name = safe_sheet_name(base, prefix=f"R{audit.sequence_no:03d}-", used=used)
            source_book = template_books.get(audit.checklist_source_workbook or "")
            source_sheet = None
            if source_book is not None and audit.checklist_source_sheet in source_book.sheetnames:
                source_sheet = source_book[audit.checklist_source_sheet]
            checklist = book.create_sheet(checklist_name)
            rows = _build_checklist(checklist, snapshot, phase, audit, source_sheet)
            report = book.create_sheet(report_name)
            report_rows = _build_audit_report(
                report,
                snapshot,
                phase,
                audit,
                checklist_name,
                rows,
                report_templates[_AUDIT_REPORT_TEMPLATE_SHEET],
            )
            infos.append((phase, _AuditSheetInfo(audit, checklist_name, report_name, report_rows)))
        directory.merge_cells("A1:E1"); directory["A1"] = "公司级质量保证过程审核目录"; directory["A1"].fill = _TITLE; directory["A1"].font = Font(name="宋体", size=18, bold=True, color="FFFFFF")
        directory.append(["阶段", "流水号", "审核对象", "检查单", "审核报告"]); _style(directory, "A2:E2", _HEADER)
        for index, (phase, info) in enumerate(infos, 3):
            directory.cell(index, 1, phase.name); directory.cell(index, 2, info.audit.sequence_no); write_literal_cell(directory.cell(index, 3), _audit_object(info.audit))
            for col, text, sheet in ((4, "打开检查单", info.checklist_sheet), (5, "打开审核报告", info.report_sheet)):
                cell = directory.cell(index, col, text); cell.hyperlink = f"#'{sheet}'!A1"; cell.style = "Hyperlink"
            _style(directory, f"A{index}:E{index}")
        for col, width in zip("ABCDE", (22, 12, 36, 16, 16), strict=True): directory.column_dimensions[col].width = width
        directory.freeze_panes = "A3"; directory.print_area = f"A1:E{len(infos)+2}"; directory.sheet_properties.pageSetUpPr.fitToPage = True; directory.page_setup.fitToWidth = 1
        book.save(output_path)
    finally:
        book.close()
        product_templates.close()
        process_templates.close()
        report_templates.close()
    return output_path

def build_trend_workbook(snapshot: ExportSnapshot, phase_paths: tuple[Path, ...], output_path: Path) -> Path:
    book = Workbook(); _set_recalculation(book); ws = book.active; ws.title = TREND_SHEET; ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:O1"); ws["A1"] = "项目质量工作趋势分析"; ws["A1"].fill = _TITLE; ws["A1"].font = Font(name="宋体", size=18, bold=True, color="FFFFFF")
    headers = ("阶段", "NC总数", "轻微", "一般", "严重", "RDM", "PP", "PMC", "CM", "MT", "MPM", "VV", "LD", "II", "工作产品")
    for col, value in enumerate(headers, 1): ws.cell(3, col, value)
    _style(ws, "A3:O3", _HEADER)
    path_by_phase = {path.name.split("-", 5)[3]: path for path in phase_paths}
    links = {}
    for phase in snapshot.phases:
        if any(a.report_scope == "phase" for a in phase.audits):
            filename = next(path.name for path in phase_paths if f"-{phase.sort_order:02d}-" in path.name)
            links[phase.id] = add_external_workbook_link(book, target=filename, sheet_names=(CALCULATION_SHEET,))
    for row, phase in enumerate(snapshot.phases, 4):
        ws.cell(row, 1, phase.name); link = links.get(phase.id)
        if link:
            ref = lambda cell: external_reference(link, CALCULATION_SHEET, cell)
            labels = ref("$A:$A")
            values = ref("$B:$B")
            for col, label in ((2, "NC总数"), (3, "轻微"), (4, "一般"), (5, "严重")):
                write_formula(
                    ws.cell(row, col),
                    Formula(f'=IFERROR(INDEX({values},MATCH("{label}",{labels},0)),0)'),
                )
            for col, category in enumerate(_CATEGORIES, 6):
                write_formula(
                    ws.cell(row, col),
                    Formula(f'=IFERROR(INDEX({values},MATCH("{category}",{labels},0)),0)'),
                )
        else:
            for col in range(2, 16): ws.cell(row, col, 0)
    total = 4 + len(snapshot.phases)
    ws.cell(total, 1, "合计")
    for col in range(2, 16): write_formula(ws.cell(total, col), Formula(f"=SUM({ws.cell(4,col).coordinate}:{ws.cell(total-1,col).coordinate})"))
    conclusion = snapshot.trend_conclusion_override or "=IF(B%d=0,\"当前项目未发现不符合项。\",\"当前项目共发现\"&B%d&\"项不符合项，其中严重\"&E%d&\"项。\")" % (total, total, total)
    ws.merge_cells(f"A{total+2}:O{total+3}"); write_formula(ws.cell(total+2, 1), Formula(conclusion)) if conclusion.startswith("=") else write_literal_cell(ws.cell(total+2, 1), conclusion)
    _style(ws, f"A{total}:O{total}", _HEADER); _style(ws, f"A{total+2}:O{total+3}")
    # The trend chart deliberately has two lines: total NC volume and severe
    # NC volume.  The other severities remain in the table for auditability.
    line = LineChart(); line.title = "阶段不符合项数量趋势图"
    line.add_data(Reference(ws, min_col=2, max_col=2, min_row=3, max_row=total-1), titles_from_data=True)
    line.add_data(Reference(ws, min_col=5, max_col=5, min_row=3, max_row=total-1), titles_from_data=True)
    line.set_categories(Reference(ws, min_col=1, min_row=4, max_row=total-1)); ws.add_chart(line, f"A{total+5}")

    distribution_start = 3
    ws.cell(distribution_start, 17, "类别")
    ws.cell(distribution_start, 18, "问题数")
    ws.cell(distribution_start, 19, "占比")
    _style(ws, f"Q{distribution_start}:S{distribution_start}", _HEADER)
    for row, (category, source_col) in enumerate(zip(_CATEGORIES, range(6, 16), strict=True), distribution_start + 1):
        ws.cell(row, 17, category)
        write_formula(ws.cell(row, 18), Formula(f"={ws.cell(total, source_col).coordinate}"))
        write_formula(ws.cell(row, 19), Formula(f"=IFERROR(R{row}/$B${total},0)"))
        ws.cell(row, 19).number_format = "0.0%"
    _style(ws, f"Q{distribution_start + 1}:S{distribution_start + len(_CATEGORIES)}")
    pie = PieChart(); pie.title = "十类不符合项分布比例图"
    pie.add_data(Reference(ws, min_col=19, min_row=distribution_start, max_row=distribution_start + len(_CATEGORIES)), titles_from_data=True)
    pie.set_categories(Reference(ws, min_col=17, min_row=distribution_start + 1, max_row=distribution_start + len(_CATEGORIES)))
    ws.add_chart(pie, f"I{total+5}")
    for col in range(1,16): ws.column_dimensions[chr(64+col)].width = 14
    for col, width in zip(("Q", "R", "S"), (14, 12, 12), strict=True): ws.column_dimensions[col].width = width
    ws.print_area = f"A1:O{total+22}"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    book.save(output_path); book.close(); return output_path

def build_complete_export_package(snapshot: ExportSnapshot, output_root: Path, build_phases) -> CompanyTrendArtifacts:
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    output_dir = output_root / f"QualityM-{snapshot.project_code}-V{snapshot.version_no}-{stamp}"
    phase_artifacts = build_phases(snapshot, output_dir)
    phase_paths = tuple(item.path for item in phase_artifacts)
    company = build_company_workbook(snapshot, output_dir / "公司级质量保证过程审核记录.xlsx")
    trend = build_trend_workbook(snapshot, phase_paths, output_dir / "项目质量工作趋势分析.xlsx")
    manifest = output_dir / "manifest.txt"; manifest.write_text("\n".join(path.name for path in (*phase_paths, *( (company,) if company else () ), trend)) + "\n", encoding="utf-8")
    return CompanyTrendArtifacts(output_dir, phase_paths, company, trend, manifest)
