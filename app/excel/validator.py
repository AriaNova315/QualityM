from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook

from app.excel.spike_builder import (
    AUDIT_REPORT_SHEET,
    CHECKLIST_SHEET,
    PHASE_REPORT_SHEET,
    TREND_SHEET,
)
from app.excel.phase_workbook_builder import (
    CALCULATION_SHEET,
    DIRECTORY_SHEET,
    PHASE_REPORT_SHEET as COMPLETE_PHASE_REPORT_SHEET,
)
from app.excel.company_trend_builder import COMPANY_DIRECTORY_SHEET, TREND_SHEET


@dataclass(frozen=True, slots=True)
class SpikeValidation:
    phase_sheets: tuple[str, ...]
    external_link_parts: tuple[str, ...]
    chart_parts: tuple[str, ...]


@dataclass(frozen=True, slots=True)
class PhaseWorkbookValidation:
    sheet_names: tuple[str, ...]
    audit_count: int
    formula_count: int
    issue_count: int


def validate_company_trend_workbooks(company_path: Path | None, trend_path: Path) -> None:
    """Structural guard for Step 9's cross-workbook package."""
    trend = load_workbook(trend_path, data_only=False, keep_links=True)
    try:
        if trend.sheetnames != [TREND_SHEET] or len(trend[TREND_SHEET]._charts) != 2:
            raise ValueError("Trend workbook is missing its sheet or charts")
        if trend.calculation.calcMode != "auto":
            raise ValueError("Trend workbook automatic calculation is disabled")
        for worksheet in trend.worksheets:
            for row in worksheet.iter_rows():
                for cell in row:
                    if cell.data_type == "f" and ("#REF!" in str(cell.value) or "{{" in str(cell.value)):
                        raise ValueError(f"Unsafe trend formula at {cell.coordinate}")
    finally:
        trend.close()
    if company_path is not None:
        company = load_workbook(company_path, data_only=False)
        try:
            if COMPANY_DIRECTORY_SHEET not in company.sheetnames:
                raise ValueError("Company workbook is missing its directory")
            directory = company[COMPANY_DIRECTORY_SHEET]
            if not directory["D3"].hyperlink or not directory["E3"].hyperlink:
                raise ValueError("Company directory links are missing")
        finally:
            company.close()


def validate_excel_linkage_spike(phase_path: Path, trend_path: Path) -> SpikeValidation:
    phase_workbook = load_workbook(phase_path, data_only=False)
    try:
        expected = (CHECKLIST_SHEET, AUDIT_REPORT_SHEET, PHASE_REPORT_SHEET)
        if tuple(phase_workbook.sheetnames) != expected:
            raise ValueError("Unexpected phase workbook sheet order")
        if phase_workbook[AUDIT_REPORT_SHEET]["C13"].data_type != "f":
            raise ValueError("Checklist-to-audit formula is missing")
        if phase_workbook[PHASE_REPORT_SHEET]["A91"].data_type != "f":
            raise ValueError("Audit-to-phase formula is missing")
        if phase_workbook[CHECKLIST_SHEET].protection.sheet:
            raise ValueError("Checklist sheet must remain editable")
    finally:
        phase_workbook.close()

    trend_workbook = load_workbook(trend_path, data_only=False, keep_links=True)
    try:
        if trend_workbook.sheetnames != [TREND_SHEET]:
            raise ValueError("Unexpected trend workbook sheets")
        if trend_workbook[TREND_SHEET]["B4"].data_type != "f":
            raise ValueError("Phase-to-trend external formula is missing")
        if not trend_workbook._external_links:
            raise ValueError("Trend workbook has no OOXML external link")
    finally:
        trend_workbook.close()

    with ZipFile(trend_path) as archive:
        names = tuple(sorted(archive.namelist()))
        external_links = tuple(name for name in names if name.startswith("xl/externalLinks/") and name.endswith(".xml"))
        chart_parts = tuple(name for name in names if name.startswith("xl/charts/") and name.endswith(".xml"))
        if not external_links or not chart_parts:
            raise ValueError("Trend workbook OOXML linkage or chart parts are missing")
    return SpikeValidation(
        phase_sheets=(CHECKLIST_SHEET, AUDIT_REPORT_SHEET, PHASE_REPORT_SHEET),
        external_link_parts=external_links,
        chart_parts=chart_parts,
    )


def validate_phase_workbook(path: Path) -> PhaseWorkbookValidation:
    workbook = load_workbook(path, data_only=False, read_only=False)
    try:
        names = tuple(workbook.sheetnames)
        if len(names) != len(set(names)) or any(len(name) > 31 for name in names):
            raise ValueError("Workbook contains duplicate or overlong sheet names")
        required = {DIRECTORY_SHEET, COMPLETE_PHASE_REPORT_SHEET, CALCULATION_SHEET}
        if not required.issubset(names):
            raise ValueError("Phase workbook is missing a required sheet")
        if workbook[CALCULATION_SHEET].sheet_state != "hidden":
            raise ValueError("Calculation sheet must be hidden")
        checklist_sheets = tuple(name for name in names if name.startswith("C") and "-" in name)
        report_sheets = tuple(name for name in names if name.startswith("R") and "-" in name)
        if len(checklist_sheets) != len(report_sheets):
            raise ValueError("Checklist and audit report counts differ")
        directory = workbook[DIRECTORY_SHEET]
        for row in range(5, 5 + len(report_sheets)):
            if not directory.cell(row, 3).hyperlink or not directory.cell(row, 4).hyperlink:
                raise ValueError("Directory hyperlink is missing")
        formula_count = 0
        for worksheet in workbook.worksheets:
            for row in worksheet.iter_rows():
                for cell in row:
                    if cell.data_type != "f":
                        continue
                    formula_count += 1
                    value = str(cell.value)
                    if "#REF!" in value or "[" in value or "]" in value:
                        raise ValueError(f"Unsafe or broken formula at {worksheet.title}!{cell.coordinate}")
                    if "{{" in value or "}}" in value:
                        raise ValueError(f"Unreplaced formula placeholder at {worksheet.title}!{cell.coordinate}")
        if report_sheets and formula_count == 0:
            raise ValueError("Audit-to-phase formulas are missing")
        if workbook.calculation.calcMode != "auto":
            raise ValueError("Workbook automatic calculation is disabled")
        calculation = workbook[CALCULATION_SHEET]
        issue_count = sum(
            isinstance(calculation.cell(row, 1).value, int)
            for row in range(2, calculation.max_row + 1)
        )
        return PhaseWorkbookValidation(
            sheet_names=names,
            audit_count=len(report_sheets),
            formula_count=formula_count,
            issue_count=issue_count,
        )
    finally:
        workbook.close()
