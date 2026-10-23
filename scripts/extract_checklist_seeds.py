from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from app.seed_content import generate_issue_templates, issue_family_counts
from app.seeds import analyze_seed_quality


DEFAULT_SOURCES = (
    (ROOT / "temp" / "工作产品检查单.xlsx", "product"),
    (ROOT / "temp" / "过程检查单.xlsx", "process"),
)
DEFAULT_OUTPUT = ROOT / "seeds" / "checklist-library.json"
DEFAULT_REPORT = ROOT / "seeds" / "checklist-library.coverage.json"
DEFAULT_CURATED_ISSUES_DIR = ROOT / "seeds" / "issue-templates"
ITEM_NO_PATTERN = re.compile(r"^(?P<number>[0-9]+)(?P<star>\*)?$")


CHECKLIST_KEYS = {
    "软件研制任务书检查单": "product.software-task-book",
    "软件开发计划检查单": "product.software-development-plan",
    "软件配置管理计划检查单": "product.configuration-management-plan",
    "软件质量保证计划检查单": "product.quality-assurance-plan",
    "软件测量与分析计划检查单": "product.measurement-analysis-plan",
    "软件需求规格说明检查单": "product.software-requirements-specification",
    "软件设计检查单": "product.software-design",
    "软件测试计划检查单": "product.software-test-plan",
    "软件测试说明检查单": "product.software-test-description",
    "软件测试报告检查单": "product.software-test-report",
    "软件用户手册检查单": "product.software-user-manual",
    "代码框架检查单": "product.code-framework",
    "编码规范检查单": "product.coding-standard",
    "检验报告检查单": "product.inspection-report",
    "检验大纲检查单": "product.inspection-outline",
    "需求开发与管理检查单": "process.requirements-development-management",
    "项目策划检查单": "process.project-planning",
    "项目监控检查单": "process.project-monitoring",
    "测量与绩效管理检查单": "process.measurement-performance-management",
    "质量保证检查单": "process.quality-assurance",
    "配置管理检查单": "process.configuration-management",
    "验证与确认检查单": "process.verification-validation",
    "实施基础检查单（中层验证）": "process.implementation-infrastructure",
    "领导作用检查单（中层验证）": "process.leadership",
    "运行维护检查单（中层验证）": "process.operation-maintenance",
}

GROUP_KEYS = {
    "需求开发与确认": "requirements-development-confirmation",
    "需求管理": "requirements-management",
    "工作产品评审过程": "work-product-review",
    "测试过程": "testing",
}


class ExtractionError(RuntimeError):
    """Raised when a source workbook does not match the known checklist grammar."""


@dataclass(frozen=True, slots=True)
class RowItem:
    row: int
    display_no: str
    content: str
    starred: bool


def load_curated_issue_templates(
    directory: Path = DEFAULT_CURATED_ISSUES_DIR,
) -> dict[str, dict[str, dict[str, Any]]]:
    """Load human-reviewed issue content without coupling it to workbook layout."""
    result: dict[str, dict[str, dict[str, Any]]] = {}
    if not directory.exists():
        return result
    for path in sorted(directory.glob("*.json")):
        try:
            payload = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError) as exc:
            raise ExtractionError(f"Invalid curated issue file: {path}") from exc
        if payload.get("schema_version") != 1:
            raise ExtractionError(f"Unsupported curated issue schema in {path}")
        checklist_key = payload.get("checklist_stable_key")
        items = payload.get("items")
        if not isinstance(checklist_key, str) or not isinstance(items, dict):
            raise ExtractionError(f"Invalid curated issue structure in {path}")
        if checklist_key in result:
            raise ExtractionError(f"Duplicate curated checklist: {checklist_key}")
        result[checklist_key] = items
    return result


def _curated_issues(
    checklist_key: str,
    group_key: str,
    item_key: str,
    starred: bool,
    item_config: dict[str, Any],
) -> list[dict[str, Any]]:
    default_severity = item_config.get("default_severity")
    expected_severities = {"严重"} if starred else {"轻微", "一般"}
    if default_severity not in expected_severities:
        raise ExtractionError(
            f"Curated severity for {checklist_key}/{item_key} conflicts with starred state"
        )
    raw_issues = item_config.get("issues")
    if not isinstance(raw_issues, list) or not 1 <= len(raw_issues) <= 4:
        raise ExtractionError(
            f"Curated issue count for {checklist_key}/{item_key} must be between 1 and 4"
        )
    issues: list[dict[str, Any]] = []
    local_keys: set[str] = set()
    for index, raw in enumerate(raw_issues):
        if not isinstance(raw, dict):
            raise ExtractionError(f"Invalid curated issue in {checklist_key}/{item_key}")
        local_key = raw.get("stable_key")
        if not isinstance(local_key, str) or not re.fullmatch(
            r"[a-z][a-z0-9]*(?:[._-][a-z0-9]+)*", local_key
        ):
            raise ExtractionError(f"Invalid curated issue key in {checklist_key}/{item_key}")
        if local_key in local_keys:
            raise ExtractionError(f"Duplicate curated issue key in {checklist_key}/{item_key}")
        local_keys.add(local_key)
        values = {
            field: raw.get(field)
            for field in ("title", "description", "corrective_suggestion")
        }
        if any(not isinstance(value, str) or not value.strip() for value in values.values()):
            raise ExtractionError(f"Blank curated issue content in {checklist_key}/{item_key}")
        severity = raw.get("default_severity", default_severity)
        if severity not in expected_severities:
            raise ExtractionError(
                f"Curated issue severity for {checklist_key}/{item_key}/{local_key} "
                "conflicts with starred state"
            )
        issues.append(
            {
                "stable_key": f"{checklist_key}.{group_key}.{item_key}.{local_key}",
                **values,
                "default_severity": severity,
                "placeholders": [],
                "is_default": index == 0,
            }
        )
    return issues


def _normalized_text(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def _item_key(display_no: str, content: str) -> str:
    digest = hashlib.sha256(content.encode("utf-8")).hexdigest()[:10]
    return f"item-{display_no}-{digest}"


def _source_sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _group_heading_rows(worksheet, first_item_row: int, last_item_row: int) -> dict[int, str]:
    headings: dict[int, str] = {}
    scan_start = max(6, first_item_row - 1)
    for row in range(scan_start, last_item_row + 1):
        first = _normalized_text(worksheet.cell(row, 1).value)
        second = _normalized_text(worksheet.cell(row, 2).value)
        if not first or second or ITEM_NO_PATTERN.fullmatch(first):
            continue
        if first.startswith(("说明", "注：", "1)", "2)")):
            continue
        if first not in GROUP_KEYS:
            raise ExtractionError(
                f"Unknown group heading at {worksheet.title}!A{row}: {first}"
            )
        headings[row] = first
    return headings


def extract_sheet(
    worksheet,
    source_path: Path,
    audit_kind: str,
    curated: dict[str, dict[str, dict[str, Any]]] | None = None,
) -> dict[str, Any]:
    sheet_title = worksheet.title
    checklist_key = CHECKLIST_KEYS.get(sheet_title)
    if checklist_key is None:
        raise ExtractionError(f"No stable checklist key configured for sheet: {sheet_title}")

    items: list[RowItem] = []
    for row in range(1, worksheet.max_row + 1):
        raw_no = _normalized_text(worksheet.cell(row, 1).value)
        match = ITEM_NO_PATTERN.fullmatch(raw_no)
        content = _normalized_text(worksheet.cell(row, 2).value)
        if match and content:
            items.append(
                RowItem(
                    row=row,
                    display_no=match.group("number"),
                    content=content,
                    starred=bool(match.group("star")),
                )
            )
    if not items:
        raise ExtractionError(f"No checklist items found in sheet: {sheet_title}")

    headings = _group_heading_rows(worksheet, items[0].row, items[-1].row)
    groups: list[dict[str, Any]] = []

    def append_group(name: str, stable_key: str, source_row: int | None, visible: bool) -> dict[str, Any]:
        group = {
            "stable_key": stable_key,
            "name": name,
            "display_heading": visible,
            "source": {"row": source_row},
            "items": [],
        }
        groups.append(group)
        return group

    current_group: dict[str, Any] | None = None
    for item in items:
        preceding = [(row, title) for row, title in headings.items() if row < item.row]
        heading = max(preceding, default=None, key=lambda value: value[0])
        if heading is None:
            if current_group is None:
                current_group = append_group("通用检查项", "default", None, False)
        else:
            row, title = heading
            stable_key = GROUP_KEYS[title]
            if current_group is None or current_group["stable_key"] != stable_key:
                current_group = append_group(title, stable_key, row, True)
        assert current_group is not None
        key = _item_key(item.display_no, item.content)
        curated_item = (curated or {}).get(checklist_key, {}).get(key)
        issues = (
            _curated_issues(
                checklist_key,
                current_group["stable_key"],
                key,
                item.starred,
                curated_item,
            )
            if curated_item is not None
            else generate_issue_templates(
                checklist_key,
                current_group["stable_key"],
                key,
                item.content,
                item.starred,
            )
        )
        current_group["items"].append(
            {
                "stable_key": key,
                "display_no": item.display_no,
                "content": item.content,
                "starred": item.starred,
                "source": {"row": item.row, "content_cell": f"B{item.row}"},
                "issue_templates": issues,
            }
        )

    curated_items = (curated or {}).get(checklist_key)
    if curated_items is not None:
        extracted_keys = {item["stable_key"] for group in groups for item in group["items"]}
        if set(curated_items) != extracted_keys:
            missing = sorted(extracted_keys - set(curated_items))
            unknown = sorted(set(curated_items) - extracted_keys)
            raise ExtractionError(
                f"Curated coverage mismatch for {checklist_key}: "
                f"missing={missing}, unknown={unknown}"
            )

    title = _normalized_text(worksheet["A1"].value) or sheet_title
    return {
        "stable_key": checklist_key,
        "name": title,
        "audit_kind": audit_kind,
        "source_workbook": source_path.name,
        "source_sheet": sheet_title,
        "source": {
            "max_row": worksheet.max_row,
            "max_column": worksheet.max_column,
            "print_area": str(worksheet.print_area) if worksheet.print_area else None,
            "freeze_panes": str(worksheet.freeze_panes) if worksheet.freeze_panes else None,
        },
        "groups": groups,
    }


def extract_library(sources=DEFAULT_SOURCES) -> tuple[dict[str, Any], dict[str, Any]]:
    checklists: list[dict[str, Any]] = []
    source_files: list[dict[str, str]] = []
    curated = load_curated_issue_templates()
    for source_path, audit_kind in sources:
        workbook = load_workbook(source_path, data_only=False, read_only=False)
        source_files.append({"filename": source_path.name, "sha256": _source_sha256(source_path)})
        for worksheet in workbook.worksheets:
            if worksheet.title == "目录":
                continue
            checklists.append(extract_sheet(worksheet, source_path, audit_kind, curated))

    payload = {"seed_version": "2.2.0", "source_files": source_files, "checklists": checklists}
    product = [item for item in checklists if item["audit_kind"] == "product"]
    process = [item for item in checklists if item["audit_kind"] == "process"]
    all_items = [
        item
        for checklist in checklists
        for group in checklist["groups"]
        for item in group["items"]
    ]
    all_groups = [group for checklist in checklists for group in checklist["groups"]]
    report = {
        "seed_version": payload["seed_version"],
        "source_files": source_files,
        "checklist_count": len(checklists),
        "product_checklist_count": len(product),
        "process_checklist_count": len(process),
        "group_count": len(all_groups),
        "visible_group_count": sum(bool(group["display_heading"]) for group in all_groups),
        "item_count": len(all_items),
        "starred_item_count": sum(bool(item["starred"]) for item in all_items),
        "issue_template_count": sum(len(item["issue_templates"]) for item in all_items),
        "minimum_issue_templates_per_item": min(
            (len(item["issue_templates"]) for item in all_items), default=0
        ),
        "maximum_issue_templates_per_item": max(
            (len(item["issue_templates"]) for item in all_items), default=0
        ),
        "template_count_distribution": {
            str(count): sum(len(item["issue_templates"]) == count for item in all_items)
            for count in range(2, 7)
        },
        "issue_family_counts": issue_family_counts(
            issue["stable_key"]
            for item in all_items
            for issue in item["issue_templates"]
        ),
        "checklists": [
            {
                "stable_key": checklist["stable_key"],
                "name": checklist["name"],
                "audit_kind": checklist["audit_kind"],
                "group_count": len(checklist["groups"]),
                "visible_group_count": sum(
                    bool(group["display_heading"]) for group in checklist["groups"]
                ),
                "item_count": sum(len(group["items"]) for group in checklist["groups"]),
                "starred_item_count": sum(
                    bool(item["starred"])
                    for group in checklist["groups"]
                    for item in group["items"]
                ),
            }
            for checklist in checklists
        ],
    }
    report["quality"] = analyze_seed_quality(payload)
    return payload, report


def write_canonical_json(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    content = json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True) + "\n"
    path.write_text(content, encoding="utf-8", newline="\n")


def main() -> None:
    parser = argparse.ArgumentParser(description="Extract QualityM checklist seed library")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--report", type=Path, default=DEFAULT_REPORT)
    args = parser.parse_args()
    payload, report = extract_library()
    write_canonical_json(args.output, payload)
    write_canonical_json(args.report, report)
    print(
        f"Extracted {report['checklist_count']} checklists, {report['item_count']} items, "
        f"{report['issue_template_count']} issue templates"
    )


if __name__ == "__main__":
    main()
